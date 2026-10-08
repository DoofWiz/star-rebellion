#!/usr/bin/env python3
"""Star Rebellion game database tooling.

The database lives in game/data/db.json. This script validates it and moves it
to and from a spreadsheet so it can be edited in Google Sheets.

    python3 tools/db/build.py validate
    python3 tools/db/build.py report
    python3 tools/db/build.py export-js
    python3 tools/db/build.py export-xlsx  [out.xlsx]
    python3 tools/db/build.py import-xlsx  in.xlsx [--dry-run]

Needs openpyxl for the xlsx commands (pip install openpyxl).
"""
import argparse
import json
import math
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
DB_PATH = ROOT / "game" / "data" / "db.json"
JS_PATH = ROOT / "game" / "data" / "db.js"

# --------------------------------------------------------------------------
# Schema. One entry per table, in the order they appear in db.json.
# Column types: id (nullable string), text (string, blank = ""), int, float,
# num (int or float), bool.
# --------------------------------------------------------------------------
KINDS_SHIP = ["starship", "drone", "structure"]
KINDS_WEAPON = ["plasma", "ballistic", "missile"]
KINDS_PILOT = ["rebel", "enemy", "drone"]
CLASSES = ["starship", "capital"]
ITEM_CATS = ["weapon", "armour", "gadget", "other", "builtin"]
ITEM_SLOTS = ["primary", "secondary", "back", "head", "body", "gadget"]
ITEM_ORIGINS = ["factory", "handmade", "scavenged"]
DAMAGE_TYPES = ["ballistic", "plasma", "laser", "explosive", "blunt"]
FACTIONS = ["hegemony", "outworlder"]
ENEMY_KINDS = ["person", "auto", "bot", "vehicle"]
OWNED_AS = ["policebot", "bruiser", "strider", "police", "dispersal", "transport"]
FIRE_MODES = ["auto", "semi", "single", "fan"]
ITEM_STATS = ["damage_min", "damage_max", "range", "attack", "shots", "damage_type"]

SCHEMA = {
    "rules": {
        "sheet": "Rules", "key": "key",
        "cols": [
            ("key", "id", 22, "Rule name. Formulas look these up by name, so do not rename."),
            ("value", "num", 10, "The number the rule uses."),
            ("note", "text", 110, "What the rule means."),
        ],
    },
    "manufacturers": {
        "sheet": "Manufacturers", "key": "id",
        "cols": [
            ("id", "id", 16, "Unique lowercase id. Other tables point at this."),
            ("name", "text", 26, "Display name."),
            ("description", "text", 90, "Flavour text from the Gear doc."),
            ("makes", "text", 24, "What they make."),
        ],
    },
    "size_scale": {
        "sheet": "Size Scale", "key": "size",
        "cols": [
            ("size", "int", 8, "Size 1 is a human, 5 is a fighter jet, 20 is a super carrier."),
            ("class", "id", 12, "starship (small freighters and below) or capital (anything bigger)."),
            ("reference", "text", 56, "Example things at this size."),
            ("approx_length_m", "num", 16, "Rough length in metres. Guidance only."),
        ],
        "auto": ["base_tn (auto)"],
    },
    "weapons": {
        "sheet": "Weapons", "key": "id",
        "cols": [
            ("id", "id", 24, "Unique lowercase id, e.g. bls-t-light-repeaters."),
            ("name", "text", 26, "Display name."),
            ("kind", "id", 11, "plasma, ballistic or missile. Drives the icon and base rules."),
            ("manufacturer", "id", 14, "Manufacturer id. Blank is fine."),
            ("damage_min", "int", 10, "Lowest damage roll."),
            ("damage_max", "int", 10, "Highest damage roll."),
            ("range", "int", 8, "Weapon range in game units."),
            ("accuracy_mod", "int", 10, "Added to the attacker's roll. Negative is harder to hit with."),
            ("ammo", "int", 8, "Rounds carried. Blank means unlimited."),
            ("crit_chance", "float", 10, "0 to 1. 0.25 = 25% chance to crit."),
            ("bypasses_shields", "bool", 12, "TRUE if shields do not absorb it."),
            ("needs_lock", "bool", 10, "TRUE if a lock-on is needed to fire."),
            ("fire_support", "id", 14, "Ground-mission fire support this weapon enables on a ship with a gunner position, e.g. door_gunner. Blank for none."),
            ("special_rules", "text", 60, "Anything the numbers do not cover."),
            ("description", "text", 50, "Flavour text."),
            ("legacy_keys", "text", 14, "Key used in the current code, so the old and new data can be matched."),
            ("notes", "text", 60, "Designer notes. Not shown to players."),
        ],
    },
    "ships": {
        "sheet": "Ships", "key": "id",
        "cols": [
            ("id", "id", 20, "Unique lowercase id. Generated; players never see it."),
            ("name", "text", 22, "Display name."),
            ("kind", "id", 10, "starship, drone (no pilot) or structure (never moves)."),
            ("manufacturer", "id", 14, "Manufacturer id. Blank is fine."),
            ("role", "text", 24, "Short role, e.g. Interceptor."),
            ("size", "int", 6, "1 to 20. See Size Scale. Bigger ships are easier to hit."),
            ("straight_max", "int", 9, "Fastest straight move."),
            ("straight_min", "int", 9, "Slowest straight move. Same as max means that is the only speed."),
            ("bank_max", "int", 9, "Fastest bank. Blank means the ship cannot bank."),
            ("bank_min", "int", 9, "Slowest bank."),
            ("turn_max", "int", 9, "Fastest hard turn. Blank means the ship cannot hard-turn."),
            ("turn_min", "int", 9, "Slowest hard turn."),
            ("shield_front", "int", 9, "Front shield value."),
            ("shield_rear", "int", 9, "Rear shield value."),
            ("armour", "int", 8, "Armour value."),
            ("hull", "int", 7, "Hull value."),
            ("weapon_slots", "int", 8, "Number of weapon slots. Any weapon fits any slot."),
            ("hard_points", "int", 8, "Cap on attachment points the ship supports. Attachments are not weapons."),
            ("default_weapon_1", "id", 22, "Weapon id in slot 1 on a stock ship. Blank is none."),
            ("default_weapon_2", "id", 22, "Weapon id in slot 2 on a stock ship. Blank is none."),
            ("extra_people", "int", 8, "People who can ride along besides the pilot."),
            ("gunner_positions", "int", 9, "Gunner positions. A weapon that enables fire support (see Weapons) needs one to be useful."),
            ("supply_drop", "bool", 10, "TRUE if, as the transport on a ground mission, it can fly a Supply Drop to the squad it carried."),
            ("fuel_per_sortie", "int", 9, "Fuel burned each mission. Blank for drones and structures."),
            ("special_rules", "text", 60, "Anything the numbers do not cover."),
            ("description", "text", 70, "Flavour text shown to the player."),
            ("legacy_keys", "text", 14, "Key(s) used in the current code, separated by |."),
            ("notes", "text", 50, "Designer notes. Not shown to players."),
        ],
        "auto": ["movement (auto)", "base_tn (auto)", "durability (auto)"],
    },
    "pilots": {
        "sheet": "Pilots", "key": "id",
        "cols": [
            ("id", "id", 20, "Unique lowercase id."),
            ("name", "text", 22, "Display name."),
            ("callsign", "text", 12, "Callsign."),
            ("kind", "id", 8, "rebel, enemy or drone (a drone's built-in flight computer)."),
            ("level", "int", 7, "1 to 20. Everyone starts at 1."),
            ("xp", "int", 6, "0 to 99. At 100 XP the pilot levels up."),
            ("initiative", "int", 9, "1 to 6. 6 moves last but shoots first."),
            ("aim", "int", 6, "Skill, 0 to 50. Raises the modifier to hit with unguided weapons."),
            ("cunning", "int", 8, "Skill, 0 to 50. Better repairs and defensive measures."),
            ("focus", "int", 7, "Skill, 0 to 50. Raises the modifier to not be hit."),
            ("presence", "int", 9, "Skill, 0 to 50. Stays Cool and avoids Panic."),
            ("ship_id", "id", 20, "Ship this pilot usually flies."),
            ("traits", "text", 22, "Rebel trait keys (the same traits rebels have, see rebel.js and rebel-exp.js), separated by |. A pair trait names its partner: friends:joss-marrek."),
            ("rank", "text", 14, "Rank title for Hegemony pilots, e.g. Commandant. Blank for rebels (their rank comes from the rebel ladder) and drones."),
            ("notes", "text", 70, "Designer notes, including how the seed values were chosen."),
        ],
        "auto": ["aim_bonus (auto)", "focus_bonus (auto)", "ship_base_tn (auto)", "target_number (auto)"],
    },
    "starting_fleet": {
        "sheet": "Starting Fleet", "key": "id",
        "cols": [
            ("id", "id", 20, "Unique lowercase id."),
            ("name", "text", 14, "The individual ship's name."),
            ("model", "id", 18, "Ship id this is an instance of."),
            ("weapon_1", "id", 22, "Weapon id in slot 1."),
            ("weapon_2", "id", 22, "Weapon id in slot 2."),
            ("notes", "text", 80, "Designer notes."),
        ],
    },
    "items": {
        "sheet": "Items", "key": "id",
        "cols": [
            ("id", "id", 13, "Unique lowercase id. Saves, loot tables and enemy loadouts use it, so do not rename."),
            ("name", "text", 26, "Display name."),
            ("sub", "text", 26, "Short line under the name, e.g. Assault rifle."),
            ("category", "id", 10, "weapon, armour, gadget or other. builtin = a weapon a unit or vehicle has built in (fists, a turret); never owned or sold."),
            ("slot", "id", 10, "primary, secondary, back, head, body or gadget. Blank: carried as mission stores, not in a slot. back is the Gear doc's back item (Heavy kit, deployables); the game has no back slot yet."),
            ("width", "int", 6, "Storeroom footprint, columns."),
            ("height", "int", 6, "Storeroom footprint, rows."),
            ("quality", "num", 7, "Auto-equip hands out the highest first. 0 = cosmetic."),
            ("manufacturer", "id", 12, "Manufacturer id. Blank if unknown or TBC."),
            ("origin", "id", 10, "factory, handmade or scavenged."),
            ("price", "int", 7, "Base Black Market price in credits. Blank: never sold, cannot be sold back."),
            ("live", "bool", 6, "TRUE once the game has rules for it. FALSE keeps it out of the market and out of every slot."),
            ("rev", "int", 5, "Lowest Revolution Level the market stocks it at."),
            ("hegemony", "bool", 9, "Hegemony issue: loot only, never sold."),
            ("drop_only", "bool", 9, "Only ever dropped by enemies, never sold."),
            ("damage_min", "int", 8, "Weapons: lowest damage roll."),
            ("damage_max", "int", 8, "Weapons: highest damage roll."),
            ("range", "int", 7, "Weapons: range in ground-map units (about 100 = one building)."),
            ("attack", "int", 7, "Weapons: added to the attack roll. Negative is harder to hit with."),
            ("shots", "int", 6, "Weapons: shots per attack."),
            ("damage_type", "id", 10, "Weapons: ballistic, plasma, laser, explosive or blunt."),
            ("one_handed", "bool", 8, "Weapons: usable with a broken arm."),
            ("jams", "bool", 6, "Weapons: can jam."),
            ("pellets", "bool", 7, "Weapons: fires a spread of pellets."),
            ("falloff", "bool", 7, "Weapons: damage drops with range."),
            ("beam", "bool", 6, "Weapons: draws as a beam."),
            ("fire_modes", "text", 10, "Weapons: fire modes the player toggles between in the attack window, separated by | (auto, semi, single, fan). Blank = no modes."),
            ("steady", "bool", 7, "Steady trait: +2 to hit a target the user held fire on (Hold)."),
            ("knockback", "bool", 9, "Knockback trait: a hit shoves the target away from the source."),
            ("stunning", "bool", 8, "Stunning trait: a hit disorients the target."),
            ("unstable", "bool", 8, "Unstable trait: a chance to blow up when used."),
            ("sundering", "bool", 9, "Weapons: Sundering trait. Damage to an armour bar is multiplied by the armour_sunder_mult rule."),
            ("piercing", "bool", 8, "Weapons: Piercing trait. The armour_pierce_frac rule's share of each hit skips the armour bar and goes to health."),
            ("heavy", "bool", 6, "Heavy trait: too heavy for the primary slot as standard (Strong rebels may; not built yet)."),
            ("deployable", "bool", 9, "Set up on the ground as an action, then used where it stands (Razorrat LMG, Riot Shield, Mining Laser)."),
            ("armour", "int", 7, "Armour: the armour bar it adds over the wearer's health in ground and boarding combat. Blank or 0 = none (cosmetic)."),
            ("shield", "int", 7, "Shield it raises when activated: a bar over armour and health, hit first. Blank = none."),
            ("icon", "id", 10, "Icon from the game's sprite sheet, e.g. gun, pistol, grenade."),
            ("description", "text", 60, "Shown to the player in the Arsenal."),
            ("notes", "text", 50, "Designer notes. Not shown to players."),
        ],
    },
    "enemies": {
        "sheet": "Enemies", "key": "id",
        "cols": [
            ("id", "id", 24, "Unique lowercase id. Mission spawns name it, so do not rename."),
            ("name", "text", 24, "Type name, e.g. Security Patrolman. Each spawn gets its own name."),
            ("faction", "id", 11, "hegemony or outworlder."),
            ("tier", "int", 5, "Revolution tier the Enemies doc puts it in (1 to 3)."),
            ("kind", "id", 8, "person, auto (a robot that fights like a person), bot (a robot vehicle that drives itself) or vehicle (needs crew; its guns are on its seats: see Vehicle Seats)."),
            ("in_doc", "bool", 7, "TRUE if the Enemies doc describes it. FALSE: the game needed it first; spec pending."),
            ("hp", "int", 6, "Health."),
            ("aim", "int", 5, "Added to attack rolls."),
            ("def", "int", 5, "Defence: the number to beat to hit them."),
            ("armour", "int", 6, "Innate armour bar (a vehicle's or Bot's plating, an Auto's shell), on top of any armour in head and body. Vehicle armour lives here, apart from ship armour. Blank = none."),
            ("shield", "int", 6, "Innate shield bar, hit before armour and health. Blank = none."),
            ("cool", "int", 6, "Nerve, 0 to 100. Blank = 55. Leaders never drop below 40."),
            ("weapon_1", "id", 12, "Item id they draw first. Blank for vehicles (their guns are on their seats)."),
            ("weapon_2", "id", 12, "Item id they also carry (a sidearm, a riot shield). Blank for none."),
            ("head", "id", 12, "Head item id. Blank for none."),
            ("body", "id", 12, "Body item id. Blank for none."),
            ("gadget", "id", 10, "Gadget item id. Blank for none."),
            ("leader", "bool", 7, "Leads the others: shoots to kill, never breaks, and the rest waver once they fall."),
            ("owned_as", "id", 10, "Robots and vehicles: the key a hacked or stolen one has at base (policebot, bruiser, strider, police, dispersal, transport)."),
            ("hack_rounds", "int", 8, "Robots only: rounds of hacking to turn it. Blank = cannot be hacked."),
            ("heavy", "bool", 6, "Hits like a truck (the Bruiser)."),
            ("big", "bool", 5, "Takes up a vehicle's footprint (the Strider)."),
            ("speed", "int", 6, "Bots and vehicles: how far one Move takes it (map units; on foot a Move is 240 and a Sprint 480). Blank for people and Autos."),
            ("credits_min", "int", 8, "Credits on the body, low end. Blank for none."),
            ("credits_max", "int", 8, "Credits on the body, high end."),
            ("art", "id", 9, "What the art kit draws: an archetype (sr-art.js ARCH) for people and robots, a vehicle (VEHICLES) for vehicles."),
            ("description", "text", 60, "From the Enemies doc where it has one."),
            ("notes", "text", 50, "Designer notes. Not shown to players."),
        ],
    },
    "vehicle_seats": {
        "sheet": "Vehicle Seats", "key": "id",
        "cols": [
            ("id", "id", 26, "Unique lowercase id, e.g. police-cruiser-drv."),
            ("vehicle", "id", 22, "Enemies id of a vehicle. Its seats are listed in this order."),
            ("seat", "id", 7, "Short key, unique within the vehicle (drv, gun, bay1...). Missions put crew in seats by this key."),
            ("name", "text", 12, "Shown to the player, e.g. Driver, Turret, Troop bay."),
            ("drives", "bool", 7, "TRUE for the one seat that moves the vehicle."),
            ("weapon", "id", 12, "Item id of the gun this seat fires. Blank: no gun."),
            ("enclosed", "bool", 8, "TRUE: whoever sits here cannot be shot (shoot the vehicle). FALSE: an exposed seat, like a turret."),
        ],
    },
    "market_weights": {
        "sheet": "Market Weights", "key": "rev",
        "cols": [
            ("rev", "int", 5, "Revolution Level the row applies from. The market uses the highest row at or below the current level."),
            ("weapon", "num", 8, "Weight for a personal weapon lot."),
            ("gadget", "num", 8, "Weight for a gadget lot."),
            ("merc", "num", 8, "Weight for a mercenary lot."),
            ("armour", "num", 8, "Weight for an armour lot."),
            ("shipwpn", "num", 8, "Weight for a ship weapon lot."),
            ("vehicle", "num", 8, "Weight for a vehicle lot."),
            ("ship", "num", 8, "Weight for a ship lot."),
            ("notes", "text", 60, "Designer notes."),
        ],
    },
    "space_enemies": {
        "sheet": "Space Enemies", "key": "id",
        "cols": [
            ("id", "id", 20, "Unique lowercase id. Space line-ups name it, so do not rename."),
            ("name", "text", 22, "Type name, e.g. Drone Monitor. Each spawn gets its own callsign."),
            ("faction", "id", 11, "hegemony or outworlder."),
            ("tier", "int", 5, "Revolution tier the Enemies doc puts it in (1 to 3)."),
            ("in_doc", "bool", 7, "TRUE if the Enemies doc describes it. FALSE: the game needed it first; spec pending."),
            ("ship", "id", 18, "Ships id it flies. Hull, shields, movement and guns come from there."),
            ("pilot", "id", 18, "Pilots id flying it: level, skills, initiative. Drones and structures use a built-in core. A spawn can name a different pilot (each cadet is their own row)."),
            ("maneuvers", "text", 10, "Extra maneuvers, separated by |: loop (K-turn), broll (barrel roll). Blank for none."),
            ("lead", "bool", 6, "The mark: flown with the instructor's AI, and the rest of the line-up's nerve breaks when it falls."),
            ("flees", "bool", 6, "Jumps out of the sector once it panics near the edge."),
            ("calls", "id", 16, "Space Enemies id it calls in once it spots a rebel ship (once per fight). Blank for none."),
            ("clamps", "bool", 7, "Fires mag-clamps that slow the target."),
            ("nerve_cap", "int", 8, "Their nerve never rises above this (green cadets: 60). Blank: no cap."),
            ("age", "int", 5, "Pilot's age on the dossier. Blank: none shown (drones, structures)."),
            ("bio", "text", 50, "Dossier line. A spawn can give its own."),
            ("description", "text", 50, "From the Enemies doc where it has one."),
            ("notes", "text", 40, "Designer notes. Not shown to players."),
        ],
    },
}
TABLE_ORDER = list(SCHEMA)
SHEET_ORDER = ["ships", "weapons", "items", "enemies", "vehicle_seats", "space_enemies", "pilots", "starting_fleet", "market_weights", "size_scale", "manufacturers", "rules"]
REQUIRED_RULES = ["tn_base", "tn_size_divisor", "tn_floor", "skill_cap", "skill_per_bonus",
                  "level_per_bonus", "level_cap", "xp_per_level", "initiative_min", "initiative_max",
                  # repair, training and the Support specialties, read by game/js/base.js (RU)
                  "repair_rate_no_workshop", "repair_rate_workshop", "repair_cost", "train_xp_rate", "prosthetic_days",
                  "support_room_xp", "support_job_xp", "support_staff_xp_mult", "niche_train_days", "niche_train_cost",
                  "treatment_patients", "treatment_min", "treatment_max", "stabilise_min", "stabilise_max", "critical_days",
                  "comms_intel", "processing_sources", "processing_bonus", "mechanic_repair_min", "mechanic_repair_max",
                  "salvage_min", "salvage_max", "maintenance_min", "maintenance_max", "handling_sources", "handling_bonus",
                  "inventory_min", "inventory_max", "tutoring_mult", "outreach_support", "outreach_bonus",
                  "mission_support_cool", "mission_support_success", "drill_time", "drill_xp_mult", "rehab_bonus",
                  "strength_days", "strength_skill", "prosthetic_supplies", "tune_days", "frames_hull", "preflight_shield",
                  "haggler", "fence_mult", "campaign_days", "symbol_bonus", "datatap_intel", "datatap_exposure"]


# --------------------------------------------------------------------------
# Loading, saving, derived numbers
# --------------------------------------------------------------------------
def load(path=DB_PATH):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def dumps(db):
    """One row per line so git diffs stay readable."""
    out = ["{"]
    parts = [' "meta": ' + json.dumps(db.get("meta", {}), ensure_ascii=False)]
    for t in TABLE_ORDER:
        rows = db.get(t, [])
        body = ",\n".join("  " + json.dumps(r, ensure_ascii=False) for r in rows)
        parts.append(' "%s": [\n%s\n ]' % (t, body))
    out.append(",\n".join(parts))
    out.append("}")
    return "\n".join(out) + "\n"


def js_text(db):
    """The same data as a plain script, so the game loads it over file:// with no fetch."""
    return ("/* Generated from db.json by tools/db/build.py export-js. Do not edit by hand. */\n"
            "window.SR_DB=" + dumps(db).strip() + ";\n")


def rules_of(db):
    return {r["key"]: r["value"] for r in db["rules"]}


def base_tn(size, rules):
    return max(rules["tn_floor"], rules["tn_base"] - math.ceil(size / rules["tn_size_divisor"]))


def skill_bonus(skill, level, rules):
    return skill // rules["skill_per_bonus"] + level // rules["level_per_bonus"]


def span(mx, mn):
    if mx is None:
        return None
    return str(mx) if mn is None or mn == mx else "%s-%s" % (mx, mn)


def movement_text(s):
    straight = span(s["straight_max"], s["straight_min"])
    bank = span(s["bank_max"], s["bank_min"])
    turn = span(s["turn_max"], s["turn_min"])
    return "%s, %s, %s" % (
        straight + " straight" if straight else "no movement",
        bank + " Bank" if bank else "no Bank",
        turn + " Turn" if turn else "no Turn")


# --------------------------------------------------------------------------
# Validation
# --------------------------------------------------------------------------
def validate(db):
    errs = []
    E = errs.append
    for t in TABLE_ORDER:
        if t not in db:
            E("missing table %s" % t)
    if errs:
        return errs
    rules = rules_of(db)
    for k in REQUIRED_RULES:
        if k not in rules:
            E("rules: missing %s" % k)
    if errs:
        return errs

    ids = {}
    for t in TABLE_ORDER:
        key = SCHEMA[t]["key"]
        seen = set()
        for i, r in enumerate(db[t]):
            v = r.get(key)
            if v in (None, ""):
                E("%s row %d: blank %s" % (t, i + 2, key))
            elif v in seen:
                E("%s: duplicate %s %r" % (t, key, v))
            seen.add(v)
            for k, typ, _, _ in SCHEMA[t]["cols"]:
                if k not in r:
                    E("%s %r: missing column %s" % (t, v, k))
        ids[t] = seen

    def ref(t, row, col, target, required=False):
        v = row.get(col)
        if v is None:
            if required:
                E("%s %r: %s is required" % (t, row.get(SCHEMA[t]["key"]), col))
        elif v not in ids[target]:
            E("%s %r: %s %r is not in %s" % (t, row.get(SCHEMA[t]["key"]), col, v, target))

    def rng(t, row, col, lo, hi, required=True):
        v = row.get(col)
        if v is None:
            if required:
                E("%s %r: %s is required" % (t, row.get(SCHEMA[t]["key"]), col))
            return
        if not (lo <= v <= hi):
            E("%s %r: %s=%s outside %s..%s" % (t, row.get(SCHEMA[t]["key"]), col, v, lo, hi))

    for r in db["size_scale"]:
        if r["class"] not in CLASSES:
            E("size_scale %s: class %r" % (r["size"], r["class"]))
    sizes = sorted(r["size"] for r in db["size_scale"])
    if sizes != list(range(1, 21)):
        E("size_scale must cover sizes 1 to 20 exactly once")

    for w in db["weapons"]:
        n = "weapons %r" % w["id"]
        if w["kind"] not in KINDS_WEAPON:
            E("%s: kind %r" % (n, w["kind"]))
        ref("weapons", w, "manufacturer", "manufacturers")
        rng("weapons", w, "damage_min", 0, 999)
        rng("weapons", w, "damage_max", 0, 999)
        if w["damage_min"] is not None and w["damage_max"] is not None and w["damage_min"] > w["damage_max"]:
            E("%s: damage_min above damage_max" % n)
        rng("weapons", w, "range", 1, 99999)
        rng("weapons", w, "crit_chance", 0, 1)
        if w["ammo"] is not None and w["ammo"] < 1:
            E("%s: ammo must be blank (unlimited) or at least 1" % n)
        for b in ("bypasses_shields", "needs_lock"):
            if not isinstance(w[b], bool):
                E("%s: %s must be TRUE or FALSE" % (n, b))
        if w["fire_support"] not in (None, "door_gunner"):
            E("%s: unknown fire_support %r (known: door_gunner)" % (n, w["fire_support"]))

    for s in db["ships"]:
        n = "ships %r" % s["id"]
        if s["kind"] not in KINDS_SHIP:
            E("%s: kind %r" % (n, s["kind"]))
        ref("ships", s, "manufacturer", "manufacturers")
        rng("ships", s, "size", 1, 20)
        for a, b in (("straight_max", "straight_min"), ("bank_max", "bank_min"), ("turn_max", "turn_min")):
            mx, mn = s[a], s[b]
            if (mx is None) != (mn is None):
                E("%s: %s and %s must both be set or both blank" % (n, a, b))
            elif mx is not None and not (1 <= mn <= mx):
                E("%s: %s=%s %s=%s (need 1 <= min <= max)" % (n, a, mx, b, mn))
        if s["straight_max"] is None:
            E("%s: every ship needs a straight speed" % n)
        for c in ("shield_front", "shield_rear", "armour", "weapon_slots", "hard_points", "extra_people", "gunner_positions"):
            rng("ships", s, c, 0, 9999)
        if not isinstance(s["supply_drop"], bool):
            E("%s: supply_drop must be TRUE or FALSE" % n)
        elif s["supply_drop"] and not s["extra_people"]:
            E("%s: supply_drop needs a ship that carries a squad (extra_people)" % n)
        rng("ships", s, "hull", 1, 9999)
        ref("ships", s, "default_weapon_1", "weapons")
        ref("ships", s, "default_weapon_2", "weapons")
        filled = sum(1 for k in ("default_weapon_1", "default_weapon_2") if s[k])
        if filled > s["weapon_slots"]:
            E("%s: %d default weapons but only %d weapon slots" % (n, filled, s["weapon_slots"]))
        if s["kind"] == "starship":
            if not s["fuel_per_sortie"] or s["fuel_per_sortie"] < 1:
                E("%s: starships need fuel_per_sortie" % n)
        elif s["fuel_per_sortie"] is not None:
            E("%s: only starships burn fuel" % n)
        if s["kind"] == "structure" and (s["bank_max"] or s["turn_max"]):
            E("%s: structures cannot bank or turn" % n)

    for p in db["pilots"]:
        n = "pilots %r" % p["id"]
        if p["kind"] not in KINDS_PILOT:
            E("%s: kind %r" % (n, p["kind"]))
        rng("pilots", p, "level", 1, rules["level_cap"])
        rng("pilots", p, "xp", 0, rules["xp_per_level"] - 1)
        rng("pilots", p, "initiative", rules["initiative_min"], rules["initiative_max"])
        for sk in ("aim", "cunning", "focus", "presence"):
            rng("pilots", p, sk, 0, rules["skill_cap"])
        ref("pilots", p, "ship_id", "ships")
        for t in [x for x in p["traits"].split("|") if x]:
            k, _, partner = t.partition(":")
            if not k or not k.replace("_", "").isalnum():
                E("%s: trait %r is not a trait key" % (n, t))
            if partner and partner not in ids["pilots"]:
                E("%s: trait %r names a partner that is not in pilots" % (n, t))
        if p["kind"] != "enemy" and p["rank"]:
            E("%s: only Hegemony pilots carry a rank title (rebels use the rebel ladder)" % n)

    ship_by_id = {s["id"]: s for s in db["ships"]}
    for f in db["starting_fleet"]:
        n = "starting_fleet %r" % f["id"]
        ref("starting_fleet", f, "model", "ships", required=True)
        ref("starting_fleet", f, "weapon_1", "weapons")
        ref("starting_fleet", f, "weapon_2", "weapons")
        m = ship_by_id.get(f["model"])
        if m:
            filled = sum(1 for k in ("weapon_1", "weapon_2") if f[k])
            if filled > m["weapon_slots"]:
                E("%s: %d weapons on a ship with %d slots" % (n, filled, m["weapon_slots"]))

    for it in db["items"]:
        n = "items %r" % it["id"]
        for col, allowed in (("category", ITEM_CATS), ("slot", ITEM_SLOTS + [None]), ("origin", ITEM_ORIGINS + [None]),
                             ("damage_type", DAMAGE_TYPES + [None])):
            if it[col] not in allowed:
                E("%s: %s %r (allowed: %s)" % (n, col, it[col], ", ".join(a for a in allowed if a)))
        ref("items", it, "manufacturer", "manufacturers")
        for b in ("live", "hegemony", "drop_only", "one_handed", "jams", "pellets", "falloff", "beam", "steady", "knockback", "stunning", "unstable", "sundering", "piercing", "heavy", "deployable"):
            if not isinstance(it[b], bool):
                E("%s: %s must be TRUE or FALSE" % (n, b))
        rng("items", it, "width", 1, 4)
        rng("items", it, "height", 1, 4)
        rng("items", it, "rev", 1, 5)
        rng("items", it, "price", 1, 99999, required=False)
        for fm in [x for x in it["fire_modes"].split("|") if x]:
            if fm not in FIRE_MODES:
                E("%s: unknown fire mode %r (known: %s)" % (n, fm, ", ".join(FIRE_MODES)))
        rng("items", it, "armour", 0, 999, required=False)
        rng("items", it, "shield", 0, 999, required=False)
        if it["armour"] and it["category"] != "armour":
            E("%s: only armour items carry an armour value" % n)
        has = [c for c in ITEM_STATS if it[c] is not None]
        if has and len(has) < len(ITEM_STATS):
            E("%s: weapon stats are all-or-nothing; missing %s" % (n, ", ".join(c for c in ITEM_STATS if it[c] is None)))
        if it["category"] == "builtin":
            if not has:
                E("%s: a builtin weapon needs its stats" % n)
            if it["slot"] or it["price"] is not None:
                E("%s: builtin weapons have no slot and no price" % n)
        if it["category"] == "weapon" and it["live"] and not has and not it["deployable"]:
            E("%s: a live weapon needs its stats (or set live FALSE); only a deployable such as the riot shield may have none" % n)
        if has:
            rng("items", it, "damage_min", 0, 999)
            rng("items", it, "damage_max", 0, 999)
            if it["damage_min"] > it["damage_max"]:
                E("%s: damage_min above damage_max" % n)
            rng("items", it, "range", 1, 9999)
            rng("items", it, "shots", 1, 20)
        if it["category"] == "weapon" and it["slot"] not in ("primary", "secondary", "back"):
            E("%s: weapons go in the primary, secondary or back slot" % n)
        if it["slot"] == "back" and not (it["heavy"] or it["deployable"]):
            E("%s: the back slot takes Heavy kit and deployables" % n)
        if it["category"] == "armour" and it["slot"] not in ("head", "body"):
            E("%s: armour goes in the head or body slot" % n)
        if not it["icon"]:
            E("%s: icon is required" % n)

    items = {it["id"]: it for it in db["items"]}
    for en in db["enemies"]:
        n = "enemies %r" % en["id"]
        for col, allowed in (("faction", FACTIONS), ("kind", ENEMY_KINDS), ("owned_as", OWNED_AS + [None])):
            if en[col] not in allowed:
                E("%s: %s %r (allowed: %s)" % (n, col, en[col], ", ".join(a for a in allowed if a)))
        for b in ("in_doc", "leader", "heavy", "big"):
            if not isinstance(en[b], bool):
                E("%s: %s must be TRUE or FALSE" % (n, b))
        rng("enemies", en, "tier", 1, 3)
        rng("enemies", en, "hp", 1, 9999)
        rng("enemies", en, "aim", -5, 20)
        rng("enemies", en, "def", 0, 30)
        rng("enemies", en, "armour", 0, 9999, required=False)
        rng("enemies", en, "shield", 0, 9999, required=False)
        rng("enemies", en, "cool", 0, 100, required=False)
        for col in ("weapon_1", "weapon_2", "head", "body", "gadget"):
            ref("enemies", en, col, "items", required=(col == "weapon_1" and en["kind"] != "vehicle"))
        if en["kind"] == "vehicle" and any(en[c] for c in ("weapon_1", "weapon_2", "head", "body", "gadget")):
            E("%s: vehicles carry no kit; their guns go on their seats (Vehicle Seats)" % n)
        w = items.get(en["weapon_1"])
        if w and w["damage_min"] is None:
            E("%s: weapon_1 %r has no combat stats" % (n, en["weapon_1"]))
        for col, slot in (("head", "head"), ("body", "body"), ("gadget", "gadget")):
            it = items.get(en[col])
            if it and it["slot"] != slot:
                E("%s: %s %r goes in the %s slot" % (n, col, en[col], it["slot"]))
        if en["kind"] == "person" and en["owned_as"]:
            E("%s: only robots and vehicles have an owned_as" % n)
        if en["kind"] != "person" and not en["owned_as"]:
            E("%s: robots and vehicles need an owned_as" % n)
        if en["kind"] not in ("auto", "bot") and en["hack_rounds"] is not None:
            E("%s: only robots can be hacked" % n)
        if en["kind"] in ("bot", "vehicle"):
            rng("enemies", en, "speed", 1, 2000)
        elif en["speed"] is not None:
            E("%s: only bots and vehicles have a speed" % n)
    owned = [en["owned_as"] for en in db["enemies"] if en["owned_as"]]
    for k in set(owned):
        if owned.count(k) > 1:
            E("enemies: owned_as %r is used twice" % k)

    kinds = {en["id"]: en["kind"] for en in db["enemies"]}
    seats = {}
    for st in db["vehicle_seats"]:
        n = "vehicle_seats %r" % st["id"]
        ref("vehicle_seats", st, "vehicle", "enemies", required=True)
        ref("vehicle_seats", st, "weapon", "items")
        if st["vehicle"] in kinds and kinds[st["vehicle"]] != "vehicle":
            E("%s: %r is not a vehicle" % (n, st["vehicle"]))
        for b in ("drives", "enclosed"):
            if not isinstance(st[b], bool):
                E("%s: %s must be TRUE or FALSE" % (n, b))
        w = items.get(st["weapon"])
        if w and w["damage_min"] is None:
            E("%s: weapon %r has no combat stats" % (n, st["weapon"]))
        seats.setdefault(st["vehicle"], []).append(st)
    for vid, k in kinds.items():
        if k != "vehicle":
            continue
        L = seats.get(vid, [])
        if sum(1 for st in L if st["drives"]) != 1:
            E("enemies %r: a vehicle needs exactly one seat that drives" % vid)
        keys = [st["seat"] for st in L]
        if len(set(keys)) != len(keys):
            E("enemies %r: seat keys must be unique" % vid)

    revs = [r["rev"] for r in db["market_weights"]]
    if 1 not in revs:
        E("market_weights: needs a row for rev 1")
    for r in db["market_weights"]:
        rng("market_weights", r, "rev", 1, 5)
        for c in ("weapon", "gadget", "merc", "armour", "shipwpn", "vehicle", "ship"):
            rng("market_weights", r, c, 0, 9999)

    ships_by = {x["id"]: x for x in db["ships"]}
    for se in db["space_enemies"]:
        n = "space_enemies %r" % se["id"]
        if se["faction"] not in FACTIONS:
            E("%s: faction %r" % (n, se["faction"]))
        rng("space_enemies", se, "tier", 1, 3)
        ref("space_enemies", se, "ship", "ships", required=True)
        ref("space_enemies", se, "pilot", "pilots", required=True)
        ref("space_enemies", se, "calls", "space_enemies")
        if se["calls"] == se["id"]:
            E("%s: cannot call itself" % n)
        for b in ("in_doc", "lead", "flees", "clamps"):
            if not isinstance(se[b], bool):
                E("%s: %s must be TRUE or FALSE" % (n, b))
        for m in [x for x in se["maneuvers"].split("|") if x]:
            if m not in ("loop", "broll"):
                E("%s: unknown maneuver %r (known: loop, broll)" % (n, m))
        sh = ships_by.get(se["ship"])
        if sh and sh["kind"] == "structure" and (se["maneuvers"] or se["flees"] or se["calls"] or se["clamps"]):
            E("%s: a structure cannot maneuver, flee, call or clamp" % n)
        rng("space_enemies", se, "age", 1, 120, required=False)
        rng("space_enemies", se, "nerve_cap", 0, 100, required=False)
        if (en["credits_min"] is None) != (en["credits_max"] is None):
            E("%s: credits_min and credits_max must both be set or both blank" % n)
        elif en["credits_min"] is not None and en["credits_min"] > en["credits_max"]:
            E("%s: credits_min above credits_max" % n)
    return errs


# --------------------------------------------------------------------------
# Report
# --------------------------------------------------------------------------
def report(db):
    rules = rules_of(db)
    lines = []
    lines.append("%-18s %4s %3s  %-34s %-7s %3s %4s %4s %3s %2s %4s" %
                 ("ship", "size", "TN", "movement", "shields", "arm", "hull", "dur", "wpn", "hp", "fuel"))
    for s in db["ships"]:
        dur = s["shield_front"] + s["shield_rear"] + s["armour"] + s["hull"]
        lines.append("%-18s %4d %3d  %-34s %3d/%-3d %3d %4d %4d %3d %2d %4s" % (
            s["id"], s["size"], base_tn(s["size"], rules), movement_text(s),
            s["shield_front"], s["shield_rear"], s["armour"], s["hull"], dur,
            s["weapon_slots"], s["hard_points"], s["fuel_per_sortie"] if s["fuel_per_sortie"] else "-"))
    lines.append("")
    lines.append("%-18s %3s %4s %-18s %4s %4s %3s" % ("pilot", "lvl", "init", "ship", "aim+", "foc+", "TN"))
    ships = {s["id"]: s for s in db["ships"]}
    for p in db["pilots"]:
        fb = skill_bonus(p["focus"], p["level"], rules)
        ab = skill_bonus(p["aim"], p["level"], rules)
        tn = base_tn(ships[p["ship_id"]]["size"], rules) + fb if p["ship_id"] else "-"
        lines.append("%-18s %3d %4d %-18s %4d %4d %3s" % (p["id"], p["level"], p["initiative"], p["ship_id"] or "-", ab, fb, tn))
    return "\n".join(lines)


# --------------------------------------------------------------------------
# xlsx export / import
# --------------------------------------------------------------------------
def _need_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        sys.exit("openpyxl is required: pip install openpyxl")


def R(key):
    """Formula snippet that reads a value from the Rules tab by its name."""
    return 'INDEX(Rules!$B$2:$B$100,MATCH("%s",Rules!$A$2:$A$100,0))' % key


GUIDE = [
    ("STAR REBELLION DATABASE", "title"),
    ("This workbook is a copy of game/data/db.json. Edit it here, download it as .xlsx and send it back; the tool turns it into db.json again. The design docs remain the source of truth for what the numbers should be.", "text"),
    ("", "text"),
    ("HOW TO EDIT", "h"),
    ("1. One tab per table. Row 1 is the header and must not be renamed, deleted or reordered. Add rows underneath.", "text"),
    ("2. Every row needs a unique id (lowercase, hyphens, e.g. sf-11-talon). Other tabs refer to it. Players never see ids, so rename freely until something points at the id.", "text"),
    ("3. Blank cell = none. A blank bank or turn means the ship cannot make that move. A blank ammo means unlimited. Blank fuel is for drones and structures.", "text"),
    ("4. Grey columns marked (auto) are formulas. They are for your eyes and are ignored when the file is read back. Do not type over them.", "text"),
    ("5. Dropdowns on kind, manufacturer, weapon and ship columns list the valid choices. Hover a header for what the column means.", "text"),
    ("6. The Rules tab holds the numbers the formulas use (TN base, skill cap and so on). Change a value there and every (auto) column updates.", "text"),
    ("", "text"),
    ("MOVEMENT PROFILES", "h"),
    ("Each move has a max and a min speed. 4-1 straight means the ship can fly straight at speed 1, 2, 3 or 4. If max and min match, that is the only speed. If both are blank the ship cannot make that move. Left and right are always the same.", "text"),
    ("", "text"),
    ("HOW TARGET NUMBER WORKS", "h"),
    ("Base TN = tn_base - ROUNDUP(size / tn_size_divisor), never below tn_floor. That is the TN against a pilot with no Focus. A pilot adds Focus bonus = FLOOR(focus / skill_per_bonus) + FLOOR(level / level_per_bonus). Range, evasive manoeuvres, cover and the rest are added on top in the game, not here.", "text"),
    ("The Pilots tab shows each pilot's TN in their usual ship (target_number). Aim uses the same bonus formula, so equal skill and level cancel out.", "text"),
    ("", "text"),
    ("TABLES", "h"),
    ("Ships: starships, drones and structures. Weapons: every ship weapon. Items: personal kit (weapons, armour, gadgets) and the built-in weapons of units and vehicles; the game draws every item from here. Enemies: the ground roster (faction, stats, what they carry, which is also what they drop), robots and vehicles included. Vehicle Seats: each vehicle's seats and the guns on them. Space Enemies: the space roster (which ship, which pilot, how it behaves). Pilots: level, initiative and the four skills; drones carry a built-in 'core' pilot. Starting Fleet: individual ships the player begins with (a stock model plus what is loaded in it). Market Weights: how likely each kind of Black Market lot is, by Revolution Level. Size Scale: sizes 1 to 20. Manufacturers: lore list. Rules: shared numbers.", "text"),
    ("", "text"),
    ("PLACEHOLDERS", "h"),
    ("Anything labelled seed, placeholder or converted in a Notes column came from the old game values rather than from a design doc, and is there to be changed.", "text"),
]


def export_xlsx(db, out):
    _need_openpyxl()
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as L
    from openpyxl.worksheet.datavalidation import DataValidation

    def F(**k):
        k.setdefault("size", 10)
        return Font(name="Arial", **k)

    head_fill = PatternFill("solid", fgColor="1F2A44")
    auto_fill = PatternFill("solid", fgColor="E7E7E7")
    auto_head = PatternFill("solid", fgColor="7A7A7A")
    wb = Workbook()
    g = wb.active
    g.title = "Guide"
    g.column_dimensions["A"].width = 130
    for i, (txt, style) in enumerate(GUIDE, start=1):
        c = g.cell(row=i, column=1, value=txt)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = F(bold=True, size=14) if style == "title" else F(bold=True) if style == "h" else F()
    g.sheet_view.showGridLines = False

    refs = {"manufacturers": "Manufacturers", "weapons": "Weapons", "ships": "Ships", "items": "Items", "enemies": "Enemies",
            "pilots": "Pilots", "space_enemies": "Space Enemies"}
    ref_cols = {
        ("weapons", "manufacturer"): "manufacturers",
        ("ships", "manufacturer"): "manufacturers",
        ("ships", "default_weapon_1"): "weapons",
        ("ships", "default_weapon_2"): "weapons",
        ("pilots", "ship_id"): "ships",
        ("starting_fleet", "model"): "ships",
        ("starting_fleet", "weapon_1"): "weapons",
        ("starting_fleet", "weapon_2"): "weapons",
        ("items", "manufacturer"): "manufacturers",
        ("enemies", "weapon_1"): "items", ("enemies", "weapon_2"): "items", ("enemies", "head"): "items",
        ("enemies", "body"): "items", ("enemies", "gadget"): "items",
        ("vehicle_seats", "vehicle"): "enemies", ("vehicle_seats", "weapon"): "items",
        ("space_enemies", "ship"): "ships", ("space_enemies", "pilot"): "pilots", ("space_enemies", "calls"): "space_enemies",
    }
    enums = {
        ("ships", "kind"): KINDS_SHIP, ("weapons", "kind"): KINDS_WEAPON,
        ("pilots", "kind"): KINDS_PILOT, ("size_scale", "class"): CLASSES,
        ("items", "category"): ITEM_CATS, ("items", "slot"): ITEM_SLOTS, ("items", "origin"): ITEM_ORIGINS,
        ("items", "damage_type"): DAMAGE_TYPES,
        ("space_enemies", "faction"): FACTIONS, ("enemies", "faction"): FACTIONS, ("enemies", "kind"): ENEMY_KINDS, ("enemies", "owned_as"): OWNED_AS,
    }
    whole = {
        ("ships", "size"): (1, 20), ("size_scale", "size"): (1, 20),
        ("pilots", "level"): (1, 20), ("pilots", "xp"): (0, 99), ("pilots", "initiative"): (1, 6),
        ("pilots", "aim"): (0, 50), ("pilots", "cunning"): (0, 50),
        ("pilots", "focus"): (0, 50), ("pilots", "presence"): (0, 50),
    }

    sheets = {}
    for t in SHEET_ORDER:
        spec = SCHEMA[t]
        ws = wb.create_sheet(spec["sheet"])
        sheets[t] = ws
        cols = spec["cols"]
        ncol = len(cols)
        col_idx = {k: i + 1 for i, (k, _, _, _) in enumerate(cols)}
        for i, (k, typ, width, note) in enumerate(cols, start=1):
            c = ws.cell(row=1, column=i, value=k)
            c.font = F(bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
            c.comment = Comment(note, "Star Rebellion DB", width=300, height=70)
            ws.column_dimensions[L(i)].width = max(width, len(k) + 2)
        for j, name in enumerate(spec.get("auto", []), start=1):
            c = ws.cell(row=1, column=ncol + j, value=name)
            c.font = F(bold=True, color="FFFFFF")
            c.fill = auto_head
            c.alignment = Alignment(vertical="center", wrap_text=True)
            c.comment = Comment("Formula. Ignored when the file is read back.", "Star Rebellion DB", width=260, height=50)
            ws.column_dimensions[L(ncol + j)].width = 34 if name.startswith("movement") else 14
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "C2" if t != "rules" and t != "size_scale" else "B2"

        for r, row in enumerate(db[t], start=2):
            for i, (k, typ, _, _) in enumerate(cols, start=1):
                v = row.get(k)
                c = ws.cell(row=r, column=i, value=v)
                c.font = F()
                if typ == "text":
                    c.alignment = Alignment(wrap_text=True, vertical="top")
                else:
                    c.alignment = Alignment(vertical="top")
        last = max(len(db[t]) + 1, 2)

        def col(k):
            return L(col_idx[k])

        def auto(j, formula, fmt=None):
            for r in range(2, last + 1):
                c = ws.cell(row=r, column=ncol + j, value=formula(r))
                c.font = F(italic=True, color="555555")
                c.fill = auto_fill
                c.alignment = Alignment(vertical="top")
                if fmt:
                    c.number_format = fmt

        if t == "ships":
            def mv(r):
                def part(mx, mn, label, none):
                    m, n = "%s%d" % (col(mx), r), "%s%d" % (col(mn), r)
                    return 'IF(%s="",%s,%s&IF(%s=%s,"","-"&%s)&"%s")' % (m, none, m, n, m, n, label)
                s = part("straight_max", "straight_min", " straight", '"no movement"')
                b = part("bank_max", "bank_min", " Bank", '"no Bank"')
                tt = part("turn_max", "turn_min", " Turn", '"no Turn"')
                return "=%s&\", \"&%s&\", \"&%s" % (s, b, tt)
            auto(1, mv)
            tn_formula = lambda r: '=IF(%s%d="","",MAX(%s,%s-ROUNDUP(%s%d/%s,0)))' % (
                col("size"), r, R("tn_floor"), R("tn_base"), col("size"), r, R("tn_size_divisor"))
            auto(2, tn_formula)
            auto(3, lambda r: "=%s%d+%s%d+%s%d+%s%d" % (col("shield_front"), r, col("shield_rear"), r, col("armour"), r, col("hull"), r))
        if t == "size_scale":
            auto(1, lambda r: '=IF(%s%d="","",MAX(%s,%s-ROUNDUP(%s%d/%s,0)))' % (
                col("size"), r, R("tn_floor"), R("tn_base"), col("size"), r, R("tn_size_divisor")))
        if t == "pilots":
            bonus = lambda skill: (lambda r: '=IF(%s%d="","",INT(%s%d/%s)+INT(%s%d/%s))' % (
                col(skill), r, col(skill), r, R("skill_per_bonus"), col("level"), r, R("level_per_bonus")))
            auto(1, bonus("aim"))
            auto(2, bonus("focus"))
            auto(3, lambda r: '=IF(%s%d="","",IFERROR(INDEX(Ships!$%s$2:$%s$500,MATCH(%s%d,Ships!$A$2:$A$500,0)),""))' % (
                col("ship_id"), r, "AA", "AA", col("ship_id"), r))
            auto(4, lambda r: '=IF(OR(%s%d="",%s%d=""),"",%s%d+%s%d)' % (
                L(ncol + 3), r, L(ncol + 2), r, L(ncol + 3), r, L(ncol + 2), r))

        # data validation
        top = 500
        for i, (k, typ, _, _) in enumerate(cols, start=1):
            rngs = "%s2:%s%d" % (L(i), L(i), top)
            if (t, k) in ref_cols:
                tbl = ref_cols[(t, k)]
                dv = DataValidation(type="list", formula1="=%s!$A$2:$A$200" % refs[tbl], allow_blank=True)
            elif (t, k) in enums:
                dv = DataValidation(type="list", formula1='"%s"' % ",".join(enums[(t, k)]), allow_blank=True)
            elif typ == "bool":
                dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
            elif (t, k) in whole:
                lo, hi = whole[(t, k)]
                dv = DataValidation(type="whole", operator="between", formula1=str(lo), formula2=str(hi), allow_blank=True)
            else:
                continue
            dv.error = "Value not allowed here. See the header note."
            dv.showErrorMessage = True
            ws.add_data_validation(dv)
            dv.add(rngs)

    # the pilots tab reads the ships base_tn column; point it at the right letter
    ships_ncol = len(SCHEMA["ships"]["cols"])
    tn_letter = L(ships_ncol + 2)
    pw = sheets["pilots"]
    for row in pw.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, str) and "Ships!$AA$2:$AA$500" in c.value:
                c.value = c.value.replace("Ships!$AA$2:$AA$500", "Ships!$%s$2:$%s$500" % (tn_letter, tn_letter))

    for ws in wb.worksheets[1:]:
        ws.sheet_view.zoomScale = 100
    wb.save(out)


def _parse(v, typ, where):
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return "" if typ == "text" else None
    if typ in ("id", "text"):
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        return str(v).strip()
    if typ == "bool":
        if isinstance(v, bool):
            return v
        s = str(v).strip().lower()
        if s in ("true", "yes", "y", "1"):
            return True
        if s in ("false", "no", "n", "0"):
            return False
        raise ValueError("%s: %r is not TRUE or FALSE" % (where, v))
    try:
        x = float(v)
    except (TypeError, ValueError):
        raise ValueError("%s: %r is not a number" % (where, v))
    if typ == "float":
        return x
    if typ == "int":
        if not x.is_integer():
            raise ValueError("%s: %r must be a whole number" % (where, v))
        return int(x)
    return int(x) if x.is_integer() else x  # num


# columns added after workbooks were already out with the designer: an import without them keeps db.json's values
ADDED_COLS = {"ships": ("supply_drop",)}


def import_xlsx(path):
    _need_openpyxl()
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=False)
    old = load()
    db = {"meta": old.get("meta", {})}
    problems = []
    for t in TABLE_ORDER:
        spec = SCHEMA[t]
        if spec["sheet"] not in wb.sheetnames:
            problems.append("missing sheet %s" % spec["sheet"])
            continue
        ws = wb[spec["sheet"]]
        headers = {}
        for c in ws[1]:
            if c.value is not None:
                headers[str(c.value).strip()] = c.column
        missing = [k for k, _, _, _ in spec["cols"] if k not in headers]
        # a column added since the workbook was exported keeps the values already in db.json
        keep = [k for k in missing if k in ADDED_COLS.get(t, ())]
        missing = [k for k in missing if k not in keep]
        if missing:
            problems.append("%s: missing columns %s" % (spec["sheet"], ", ".join(missing)))
            continue
        if keep:
            print("%s: no %s column in the workbook (added since it was exported); keeping the values in db.json"
                  % (spec["sheet"], ", ".join(keep)))
        prev = {r.get(spec["key"]): r for r in old.get(t, [])}
        rows = []
        for r in range(2, ws.max_row + 1):
            raw = {k: ws.cell(row=r, column=headers[k]).value for k, _, _, _ in spec["cols"] if k in headers}
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in raw.values()):
                continue
            row = {}
            for k, typ, _, _ in spec["cols"]:
                if k in keep:
                    kv = raw.get(spec["key"])
                    row[k] = (prev.get(kv) or {}).get(k, False if typ == "bool" else None)
                    continue
                try:
                    row[k] = _parse(raw[k], typ, "%s row %d column %s" % (spec["sheet"], r, k))
                except ValueError as e:
                    problems.append(str(e))
            rows.append(row)
        db[t] = rows
    return db, problems


def diff(old, new):
    out = []
    for t in TABLE_ORDER:
        key = SCHEMA[t]["key"]
        a = {r[key]: r for r in old.get(t, [])}
        b = {r[key]: r for r in new.get(t, [])}
        for k in b:
            if k not in a:
                out.append("  + %s %s" % (t, k))
        for k in a:
            if k not in b:
                out.append("  - %s %s" % (t, k))
        for k in b:
            if k in a and a[k] != b[k]:
                changed = [c for c in b[k] if a[k].get(c) != b[k][c]]
                out.append("  ~ %s %s: %s" % (t, k, ", ".join("%s %r -> %r" % (c, a[k].get(c), b[k][c]) for c in changed)))
    return out


# --------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("validate")
    sub.add_parser("report")
    sub.add_parser("export-js")
    e = sub.add_parser("export-xlsx")
    e.add_argument("out", nargs="?", default="star-rebellion-db.xlsx")
    i = sub.add_parser("import-xlsx")
    i.add_argument("path")
    i.add_argument("--dry-run", action="store_true", help="show what would change without writing db.json")
    a = ap.parse_args()

    if a.cmd in ("validate", "report", "export-xlsx", "export-js"):
        db = load()
        errs = validate(db)
        if errs:
            print("db.json has %d problem(s):" % len(errs))
            for x in errs:
                print("  -", x)
            sys.exit(1)
        if a.cmd == "validate":
            counts = ", ".join("%s %d" % (t, len(db[t])) for t in TABLE_ORDER)
            print("db.json OK (%s)" % counts)
            if not JS_PATH.exists() or JS_PATH.read_text(encoding="utf-8") != js_text(db):
                print("game/data/db.js is out of date. Run: python3 tools/db/build.py export-js")
                sys.exit(1)
            print("db.js in sync")
        elif a.cmd == "export-js":
            JS_PATH.write_text(js_text(db), encoding="utf-8")
            print("wrote", JS_PATH)
        elif a.cmd == "report":
            print(report(db))
        else:
            export_xlsx(db, a.out)
            print("wrote", a.out)
    else:
        new, problems = import_xlsx(a.path)
        if not problems:
            problems = validate(new)
        if problems:
            print("Not imported. %d problem(s):" % len(problems))
            for x in problems:
                print("  -", x)
            sys.exit(1)
        changes = diff(load(), new)
        print("%d change(s):" % len(changes))
        for c in changes:
            print(c)
        if a.dry_run:
            print("(dry run, db.json untouched)")
        elif changes:
            DB_PATH.write_text(dumps(new), encoding="utf-8")
            JS_PATH.write_text(js_text(new), encoding="utf-8")
            print("wrote", DB_PATH, "and", JS_PATH)


if __name__ == "__main__":
    main()
