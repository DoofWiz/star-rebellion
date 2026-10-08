#!/usr/bin/env python3
"""Star Rebellion text tool: every line of player-facing text in one spreadsheet.

The game's dialog, briefings, tooltips and labels are written straight into the JavaScript (and the page HTML).
This tool finds them, puts them in a workbook you can edit in Google Sheets, and writes your edits back into the code.

    python3 tools/text/text.py export  [out.xlsx]            # needs: pip install openpyxl
    python3 tools/text/text.py import  in.xlsx [--dry-run]   # shows or applies your edits
    python3 tools/text/text.py find    "words"               # where does this line live?
    python3 tools/text/text.py report                        # how many lines per tab and kind

A line that is built from pieces shows the game's own values as {placeholders}: "{t.first} is down!". Keep every
placeholder, in the same order, and change the words around them. See docs/TEXT.md.
"""
import argparse
import hashlib
import os
import re
import shutil
import subprocess
import sys
import tempfile
import html
from html.parser import HTMLParser

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))

# Files scanned, in tab order. Generated files (sr-theme.js, db.js) are left out: edits there would be overwritten.
SOURCES = [
    ("Missions", ["game/js/mission-text.js"]),
    ("Prologue", ["game/js/prologue.js"]),
    ("Base", ["game/js/base.js"]),
    ("Ground", ["game/js/ground.js"]),
    ("Space", ["game/js/space.js"]),
    ("Rebels", ["game/js/rebel.js", "game/js/rebel-exp.js", "game/js/rebel-injury.js", "game/js/rebel-rest.js", "game/js/support.js"]),
    ("Title", ["game/js/intro.js"]),
    ("Shared", ["game/js/core.js", "game/js/items.js", "game/js/enemies.js", "game/js/data.js", "game/ui/sr-hud.js"]),
    ("Page", ["game/index.html"]),
]

# --------------------------------------------------------------------------
# JavaScript lexer: just enough to find every string and template literal reliably
# --------------------------------------------------------------------------
PUNCT = sorted("""
>>>= ... === !== **= <<= >>= >>> &&= ||= ??= => == != <= >= && || ?? ?. ++ -- += -= *= /= %= &= |= ^= ** << >>
{ } ( ) [ ] ; , < > + - * / % & | ^ ! ~ ? : = . @ #
""".split(), key=len, reverse=True)
REGEX_AFTER_KW = {"return", "typeof", "case", "do", "else", "in", "of", "new", "delete", "void", "throw", "instanceof",
                  "yield", "await"}
ID_START = re.compile(r"[A-Za-z_$À-￿]")
ID_CHARS = re.compile(r"[A-Za-z0-9_$À-￿]*")
NUM = re.compile(r"0[xX][0-9a-fA-F_]+n?|0[bB][01_]+n?|0[oO][0-7_]+n?|(?:\d[\d_]*\.?[\d_]*|\.\d[\d_]*)(?:[eE][+-]?\d+)?n?")


class Tok:
    __slots__ = ("t", "s", "e", "v", "rs", "re_", "nl")

    def __init__(self, t, s, e, v=None, rs=None, re_=None):
        self.t, self.s, self.e, self.v, self.rs, self.re_ = t, s, e, v, rs, re_
        self.nl = False  # a line break comes before this token

    def __repr__(self):
        return "%s(%r)" % (self.t, self.v)


class LexError(Exception):
    pass


ESC = {"n": "\n", "t": "\t", "r": "\r", "b": "\b", "f": "\f", "v": "\v", "0": "\0"}


def cook(raw):
    """The value of a string or template body, escapes resolved."""
    out, i, n = [], 0, len(raw)
    while i < n:
        c = raw[i]
        if c != "\\":
            out.append(c)
            i += 1
            continue
        i += 1
        if i >= n:
            break
        c = raw[i]
        if c in ESC and not (c == "0" and i + 1 < n and raw[i + 1].isdigit()):
            out.append(ESC[c])
            i += 1
        elif c == "x":
            out.append(chr(int(raw[i + 1:i + 3], 16)))
            i += 3
        elif c == "u":
            if raw[i + 1] == "{":
                j = raw.index("}", i)
                out.append(chr(int(raw[i + 2:j], 16)))
                i = j + 1
            else:
                out.append(chr(int(raw[i + 1:i + 5], 16)))
                i += 5
        elif c == "\r":
            i += 2 if raw[i + 1:i + 2] == "\n" else 1
        elif c in "\n  ":
            i += 1
        else:
            out.append(c)
            i += 1
    # surrogate pairs written as two \u escapes
    return "".join(out).encode("utf-16", "surrogatepass").decode("utf-16")


def lex(src):
    toks, i, n = [], 0, len(src)
    braces = []  # 'b' for a plain brace, 't' for a template substitution
    nl = False

    def prev_sig():
        return toks[-1] if toks else None

    def regex_ok():
        p = prev_sig()
        if p is None:
            return True
        if p.t in ("num", "str", "tpl", "tpl_tail", "regex"):
            return False
        if p.t == "id":
            return p.v in REGEX_AFTER_KW
        return p.v not in (")", "]", "}")

    def read_tpl(i):
        """Read template text from i up to `${` or the closing backtick. Returns (end of raw, closed?, next index)."""
        j = i
        while j < n:
            c = src[j]
            if c == "\\":
                j += 2
            elif c == "`":
                return j, True, j + 1
            elif c == "$" and src[j + 1:j + 2] == "{":
                return j, False, j + 2
            else:
                j += 1
        raise LexError("unterminated template at %d" % i)

    while i < n:
        c = src[i]
        if c in " \t﻿ ":
            i += 1
            continue
        if c in "\n\r  ":
            nl = True
            i += 1
            continue
        if src.startswith("//", i):
            j = src.find("\n", i)
            i = n if j < 0 else j
            continue
        if src.startswith("/*", i):
            j = src.find("*/", i + 2)
            if j < 0:
                raise LexError("unterminated comment at %d" % i)
            if "\n" in src[i:j]:
                nl = True
            i = j + 2
            continue
        s = i
        if c in "'\"":
            j = i + 1
            while j < n and src[j] != c:
                if src[j] == "\\":
                    j += 1
                elif src[j] == "\n":
                    raise LexError("unterminated string at %d" % i)
                j += 1
            if j >= n:
                raise LexError("unterminated string at %d" % i)
            tok = Tok("str", s, j + 1, cook(src[i + 1:j]), i + 1, j)
            i = j + 1
        elif c == "`":
            end, closed, nxt = read_tpl(i + 1)
            tok = Tok("tpl" if closed else "tpl_head", s, nxt, cook(src[i + 1:end]), i + 1, end)
            if not closed:
                braces.append("t")
            i = nxt
        elif c == "}" and braces and braces[-1] == "t":
            braces.pop()
            end, closed, nxt = read_tpl(i + 1)
            tok = Tok("tpl_tail" if closed else "tpl_mid", s, nxt, cook(src[i + 1:end]), i + 1, end)
            if not closed:
                braces.append("t")
            i = nxt
        elif ID_START.match(c) or c == "\\":
            m = ID_CHARS.match(src, i + 1)
            tok = Tok("id", s, m.end(), src[s:m.end()])
            i = m.end()
        elif c.isdigit() or (c == "." and src[i + 1:i + 2].isdigit()):
            m = NUM.match(src, i)
            tok = Tok("num", s, m.end(), src[s:m.end()])
            i = m.end()
        elif c == "/" and regex_ok():
            j, cls = i + 1, False
            while j < n:
                d = src[j]
                if d == "\\":
                    j += 2
                    continue
                if d == "\n":
                    raise LexError("unterminated regex at %d" % i)
                if d == "[":
                    cls = True
                elif d == "]":
                    cls = False
                elif d == "/" and not cls:
                    break
                j += 1
            m = ID_CHARS.match(src, j + 1)
            tok = Tok("regex", s, m.end(), src[s:m.end()])
            i = m.end()
        else:
            for p in PUNCT:
                if src.startswith(p, i):
                    break
            else:
                raise LexError("unexpected %r at %d" % (c, i))
            if p == "{":
                braces.append("b")
            elif p == "}" and braces:
                braces.pop()
            tok = Tok("p", s, i + len(p), p)
            i += len(p)
        tok.nl = nl
        nl = False
        toks.append(tok)
    return toks


# --------------------------------------------------------------------------
# From literals to lines: group a + chain into one line, name where it lives
# --------------------------------------------------------------------------
BOUND_P = set("( [ { , ; ? : = += -= *= /= %= **= <<= >>= >>>= &= |= ^= &&= ||= ??= => == === != !== < > <= >= && || ?? "
              "& | ^ << >> >>> ... ) ] }".split())
BOUND_KW = {"return", "case", "throw", "in", "of", "instanceof", "else", "do", "yield", "default", "const", "let", "var",
            "export", "import"}
ENDERS = {"num", "str", "tpl", "tpl_tail", "regex"}
UNARY_KW = {"typeof", "void", "delete", "new", "await", "return", "case", "throw", "in", "of", "instanceof", "else", "do",
            "yield"}
# a string handed to one of these is a name, a selector or a key, not text
CODE_CALLS = {"getElementById", "querySelector", "querySelectorAll", "addEventListener", "removeEventListener", "add",
              "remove", "toggle", "contains", "setAttribute", "getAttribute", "removeAttribute", "hasAttribute",
              "createElement", "createElementNS", "closest", "matches", "setProperty", "getPropertyValue", "getItem",
              "setItem", "removeItem", "indexOf", "includes", "startsWith", "endsWith", "split", "join", "replace",
              "replaceAll", "lastIndexOf", "getContext", "require", "dispatchEvent", "Event", "CustomEvent", "has",
              "get", "set", "delete", "fetch", "postMessage", "ico", "icon", "avatar", "playSound", "sfx", "tone",
              "createLinearGradient", "createRadialGradient", "rgba", "addColorStop", "test", "match", "RegExp",
              "toLocaleString", "Intl", "NumberFormat"}
COMPARE = {"==", "===", "!=", "!=="}


def is_ender(t):
    return t.t in ENDERS or (t.t == "id" and t.v not in UNARY_KW) or (t.t == "p" and t.v in (")", "]", "++", "--"))


class Scan:
    """One source file's literals, grouped into lines."""

    def __init__(self, src, base=0, line_of=None):
        self.src = src
        self.toks = lex(src)
        self.base = base
        self.match = {}
        stack = []
        for k, t in enumerate(self.toks):
            if t.t == "p" and t.v in "([{":
                stack.append(k)
            elif t.t == "tpl_head":
                stack.append(k)
            elif t.t == "p" and t.v in ")]}" or t.t == "tpl_tail":
                if stack:
                    o = stack.pop()
                    self.match[o], self.match[k] = k, o
            elif t.t == "tpl_mid":
                pass
        self.where = self._contexts()

    # ---- context ("where") --------------------------------------------------
    def _contexts(self):
        T, out, stack = self.toks, {}, []
        fname = None

        def name_before(k):
            """Name assigned at `NAME =`/`a.b.NAME =` ending just before token k (k is the `=`)."""
            j = k - 1
            if j >= 0 and T[j].t == "id":
                parts = [T[j].v]
                while j >= 2 and T[j - 1].v in (".", "?.") and T[j - 2].t == "id":
                    parts.insert(0, T[j - 2].v)
                    j -= 2
                return ".".join(parts[-2:])
            return None

        def fn_name(k):
            """If the `{` at k opens a function body, its name."""
            j = k - 1
            if j < 0:
                return None
            if T[j].v == "=>":
                j -= 1
                if T[j].v == ")" and j in self.match:
                    j = self.match[j]
                j -= 1
                if j >= 0 and T[j].v in ("=", ":"):
                    return name_before(j) if T[j].v == "=" else (T[j - 1].v if T[j - 1].t in ("id", "str") else None)
                return "(arrow)"
            if T[j].v == ")" and j in self.match:
                o = self.match[j]
                if o >= 1 and T[o - 1].t == "id" and T[o - 1].v not in ("if", "for", "while", "switch", "catch", "with"):
                    if T[o - 1].v == "function":
                        if o >= 2 and T[o - 2].v in ("=", ":"):
                            return name_before(o - 2) if T[o - 2].v == "=" else T[o - 3].v
                        return "(function)"
                    if o >= 2 and T[o - 2].v == "function":
                        return T[o - 1].v
                    return T[o - 1].v  # method shorthand
            return None

        for k, t in enumerate(T):
            if t.t == "p" and t.v in "([{" or t.t == "tpl_head":
                prev = T[k - 1] if k else None
                fr = {"o": t.v if t.t == "p" else "`", "idx": 0, "key": None, "kind": None, "name": None}
                if t.v == "(" and prev is not None and prev.t == "id" and prev.v not in UNARY_KW and prev.v not in (
                        "if", "for", "while", "switch", "catch", "function", "with"):
                    fr["kind"], fr["name"] = "call", prev.v
                elif t.v == "[" and (prev is None or not is_ender(prev)):
                    fr["kind"] = "array"
                elif t.v == "{":
                    fn = fn_name(k)
                    if fn:
                        fr["kind"], fr["name"] = "fn", fn
                    elif prev is not None and (prev.v in ("=", "(", ",", ":", "[", "?", "||", "&&", "??") or prev.v == "return"):
                        fr["kind"] = "obj"
                    else:
                        fr["kind"] = "block"
                if fr["kind"] in ("obj", "array") and prev is not None and prev.v == "=":
                    fr["name"] = name_before(k - 1)
                stack.append(fr)
                continue
            if t.t == "p" and t.v in ")]}" or t.t == "tpl_tail":
                if stack:
                    stack.pop()
                continue
            if not stack:
                if t.t in ("str", "tpl", "tpl_head"):
                    out[k] = ""
                continue
            top = stack[-1]
            if t.t == "p" and t.v == ",":
                top["idx"] += 1
                top["key"] = None
            elif top["kind"] == "obj" and k + 1 < len(T) and T[k + 1].v == ":" and T[k - 1].v in ("{", ",") and t.t in ("id", "str", "num"):
                top["key"] = t.v
            if t.t in ("str", "tpl", "tpl_head"):
                out[k] = self._path(stack)
        return out

    @staticmethod
    def _path(stack):
        """Where a literal sits, e.g. CHAINS.parity_towers.steps[1].text[0] or say() in onDown()."""
        segs, fn = [], None
        for fr in reversed(stack):
            if fr["kind"] == "fn":
                fn = fr["name"]
                break
            if fr["kind"] == "obj":
                if fr["key"] is not None:
                    segs.insert(0, "." + str(fr["key"]))
                if fr["name"]:
                    segs.insert(0, fr["name"])
                    break
            elif fr["kind"] == "array":
                segs.insert(0, "[%d]" % fr["idx"])
                if fr["name"]:
                    segs.insert(0, fr["name"])
                    break
            elif fr["kind"] == "call":
                segs.insert(0, " %s() " % fr["name"])
        # an index into an unnamed list, or a key straight after a call, reads as noise: keep the names
        out = []
        for i, g in enumerate(segs):
            if g.startswith("[") and (not out or out[-1].endswith(") ")) and (i + 1 < len(segs) and segs[i + 1].startswith(" ")):
                continue
            out.append(g)
        p = re.sub(r"\s+", " ", "".join(out)).strip()
        p = re.sub(r"\) \.", ").", p)
        if fn and fn not in ("(arrow)", "(function)"):
            p = (p + " in " if p else "") + fn + "()"
        return p.strip()

    # ---- grouping ------------------------------------------------------------
    def _bound_back(self, k):
        T, j = self.toks, k - 1
        while j >= 0:
            t = T[j]
            if t.t == "p" and t.v in (")", "]") and j in self.match:
                if T[j + 1].nl and is_ender(t) and T[j + 1].t != "p":
                    return j + 1
                j = self.match[j] - 1
                continue
            if t.t == "tpl_tail" and j in self.match:
                j = self.match[j] - 1
                continue
            if t.t in ("tpl_head", "tpl_mid") or (t.t == "p" and t.v in BOUND_P) or (t.t == "id" and t.v in BOUND_KW):
                return j + 1
            if T[j + 1].nl and is_ender(t) and (T[j + 1].t != "p"):
                return j + 1  # automatic semicolon
            j -= 1
        return 0

    def _bound_fwd(self, k):
        T, j, n = self.toks, k, len(self.toks)
        while j < n:
            t = T[j]
            if j > k and t.nl and is_ender(T[j - 1]) and t.t != "p":
                return j
            if t.t == "p" and t.v in ("(", "[") and j in self.match:
                j = self.match[j] + 1
                continue
            if t.t == "tpl_head" and j in self.match:
                j = self.match[j] + 1
                continue
            if j > k and (t.t in ("tpl_mid", "tpl_tail") or (t.t == "p" and t.v in BOUND_P) or (t.t == "id" and t.v in BOUND_KW)):
                return j
            j += 1
        return n

    def lines(self):
        """Every literal, grouped into lines. Each line: dict(pieces, first tok index, range)."""
        T, done, out = self.toks, set(), []
        for k, t in enumerate(T):
            if t.t not in ("str", "tpl", "tpl_head") or k in done:
                continue
            s, e = self._bound_back(k), self._bound_fwd(k)
            pieces, cur, j, bad = [], s, s, False
            while j < e:
                u = T[j]
                if u.t == "p" and u.v in ("(", "[") or u.t == "tpl_head":
                    j = self.match.get(j, j) + 1
                    continue
                if u.t == "p" and u.v in ("+", "-") and j > s and is_ender(T[j - 1]):
                    if u.v == "-":
                        bad = True
                        break
                    pieces.append((cur, j))
                    cur = j + 1
                j += 1
            pieces.append((cur, e))
            if bad:
                pieces = [(k, k + 1)] if t.t == "str" or t.t == "tpl" else [(k, self.match.get(k, k) + 1)]
                s, e = pieces[0]
            segs = []
            for a, b in pieces:
                if b == a + 1 and T[a].t in ("str", "tpl"):
                    segs.append(("L", a))
                elif T[a].t == "tpl_head" and self.match.get(a) == b - 1:
                    segs.append(("T", a))
                else:
                    segs.append(("E", a, b))
            # leading expressions add up before any text joins in: they are one placeholder
            first = next((i for i, g in enumerate(segs) if g[0] != "E"), None)
            if first is None:
                continue
            if first > 1:
                segs = [("E", segs[0][1], segs[first - 1][2], True)] + segs[first:]
            for g in segs:
                if g[0] == "L":
                    done.add(g[1])
                elif g[0] == "T":
                    done.add(g[1])
            out.append({"segs": segs, "k": k, "s": s, "e": e})
        return out

    def code(self, a, b):
        """Source of tokens a..b-1, whitespace collapsed."""
        if b <= a:
            return ""
        return re.sub(r"\s+", " ", self.src[self.toks[a].s:self.toks[b - 1].e]).strip()

    def flatten(self, line):
        """A line as items: ('t', text, slot) and ('x', code, where, raw code, how it sits).
        Text slots: ('s', tok) a string, ('p', tok) a piece of a template. Code sits as an operand of the + chain
        ('op'), as values added up before any text ('sum'), or inside a template's ${} ('sub')."""
        T, items = self.toks, []
        for g in line["segs"]:
            if g[0] == "L":
                items.append(("t", T[g[1]].v, ("s", g[1])))
            elif g[0] == "T":
                a, b = g[1], self.match[g[1]]
                parts = [a] + [m for m in range(a + 1, b) if T[m].t == "tpl_mid" and self._owner(m) == a] + [b]
                for i, m in enumerate(parts):
                    items.append(("t", T[m].v, ("p", m)))
                    if i + 1 < len(parts):
                        items.append(("x", self.code(m + 1, parts[i + 1]), None, self.raw(m + 1, parts[i + 1]), "sub"))
            else:
                items.append(("x", self.code(g[1], g[2]), ("e", g[1], g[2], len(g) > 3), self.raw(g[1], g[2]),
                              "sum" if len(g) > 3 else "op"))
        return items

    def raw(self, a, b):
        return self.src[self.toks[a].s:self.toks[b - 1].e] if b > a else ""

    def extent(self, line):
        """Source offsets of the whole line: from its first piece to its last."""
        segs = line["segs"]
        a = segs[0][1]
        g = segs[-1]
        b = g[1] if g[0] == "L" else self.match[g[1]] if g[0] == "T" else g[2] - 1
        return self.toks[a].s, self.toks[b].e

    def _owner(self, m):
        """The template head that a tpl_mid at m belongs to."""
        depth, j = 0, m - 1
        while j >= 0:
            t = self.toks[j]
            if t.t == "tpl_tail":
                depth += 1
            elif t.t == "tpl_head":
                if depth == 0:
                    return j
                depth -= 1
            j -= 1
        return None


SIMPLE = re.compile(r"[A-Za-z_$][\w$]*(\.[A-Za-z_$][\w$]*)*(\(\))?")


SLOT = re.compile(r"\{[A-Za-z_][\w.]*\}")   # a slot the game fills in, written into the text itself: {target}


def tokens(items):
    """What each placeholder shows as: short code as itself ({t.first}), anything longer as a number ({1})."""
    out, n = [], 0
    for x in items:
        if x[0] != "x":
            continue
        if len(x[1]) <= 28 and SIMPLE.fullmatch(x[1]) and not x[1].isdigit():
            out.append(x[1])
        else:
            n += 1
            out.append(str(n))
    return out


def marks(items):
    """Every {placeholder} the line shows, in order: (shown as, 'x' for code or 't' for a slot inside the text)."""
    toks, out, i = tokens(items), [], 0
    for x in items:
        if x[0] == "t":
            out += [(m.group(0)[1:-1], "t") for m in SLOT.finditer(x[1])]
        else:
            out.append((toks[i], "x"))
            i += 1
    return out


def show(items):
    toks, out, i = tokens(items), [], 0
    for x in items:
        if x[0] == "t":
            out.append(x[1])
        else:
            out.append("{" + toks[i] + "}")
            i += 1
    return "".join(out)


def legend(items):
    toks = tokens(items)
    codes = [x[1] for x in items if x[0] == "x"]
    return "\n".join("{%s} = %s" % (t, c) for t, c in zip(toks, codes) if t.isdigit())


# --------------------------------------------------------------------------
# What counts as player-facing text
# --------------------------------------------------------------------------
TAG = re.compile(r"<[^>]*>")
CSSISH = re.compile(r"(^|[\s;{])[-a-z]+\s*:\s*[^;:]*;|\b\d+(px|em|rem|vh|vw|ms|deg)\b|rgba?\(|var\(--|^[#.][\w-]|=>|function\s*\(|^\w+\(|"
                    r"^[\w-]+=|https?:|\.(png|svg|jpg|js|css|json|html)\b|^[\w./-]+/[\w./-]*$")
WORD = re.compile(r"[A-Za-zÀ-ɏ][A-Za-zÀ-ɏ’'\-]*")
DIALOG_WHERE = re.compile(r"line|quip|bark|say|greet|pitch|dialog|chat|speech|taunt|shout|comms|ST_|reply", re.I)


# lists of words that are not lines: random rebel names
SKIP = [("rebel.js", r"(FIRST|LAST)\["), ("", r"$^")]


def visible(text):
    t = TAG.sub(" ", text)
    t = re.sub(r"&[a-z]+;|&#\d+;", " ", t)
    return t


def code_words(scan):
    """Values the code uses as names: compared against, switched on, used as keys, looked up in a list.
    A line that is exactly one of these (a role such as 'Soldier') is an id that happens to be shown, not text."""
    T, out = scan.toks, set()
    for k, t in enumerate(T):
        prev = T[k - 1] if k else None
        nxt = T[k + 1] if k + 1 < len(T) else None
        pv, nv = prev.v if prev else None, nxt.v if nxt else None
        if t.t in ("str", "tpl"):
            if (pv in COMPARE or nv in COMPARE or pv == "case" or nv == "in" or (nv == ":" and pv in ("{", ",")) or
                    (pv == "[" and k >= 2 and is_ender(T[k - 2])) or
                    (pv == "(" and k >= 2 and T[k - 2].t == "id" and T[k - 2].v in CODE_CALLS)):
                out.add(t.v)
        elif t.t == "id" and ((nv == ":" and pv in ("{", ",")) or pv in (".", "?.")):
            out.add(t.v)
        elif t.t == "p" and t.v == "[" and k in scan.match:
            m = scan.match[k]
            if m + 2 < len(T) and T[m + 1].v in (".", "?.") and T[m + 2].v in ("includes", "indexOf"):
                out.update(T[j].v for j in range(k + 1, m) if T[j].t == "str")
    return out


def classify(scan, line, items, where, src, names=frozenset()):
    """None if this is code, else 'Dialog', 'Text' or 'Label'."""
    T = scan.toks
    texts = "".join(x[1] for x in items if x[0] == "t")
    if texts == "use strict":
        return None
    if len(items) == 1 and texts in names:
        return None
    if "{" in SLOT.sub("", texts) or "}" in SLOT.sub("", texts):
        return None
    vis = visible(texts).strip()
    words = WORD.findall(vis)
    if not words:
        return None
    s, e = line["s"], line["e"]
    before = T[s - 1] if s > 0 else None
    after = T[e] if e < len(T) else None
    if before is not None and (before.v in COMPARE or before.v == "case" or before.v == "in"):
        return None
    if after is not None and (after.v in COMPARE or after.v == "in" or (after.v == ":" and before is not None and before.v in ("{", ","))):
        return None
    if after is not None and after.v in (".", "?.", "["):
        return None  # used as a value: 'a|b'.split('|')
    if before is not None and before.v == "[" and s >= 2 and is_ender(T[s - 2]):
        return None  # obj['key']
    if before is not None and before.v in ("(", ",") and s - 1 in scan.match:
        pass
    # first argument of a call that takes names, not words
    opener = None
    j = s - 1
    if before is not None and before.v == "(":
        opener = j
    if opener is not None and opener >= 1 and T[opener - 1].t == "id" and T[opener - 1].v in CODE_CALLS:
        return None
    # what the text looks like
    raw_text = texts.strip()
    if CSSISH.search(raw_text) and not re.search(r"[.!?…]\s*$|[a-z]{3,} [a-z]{3,} [a-z]{3,}", vis):
        return None
    tagged = TAG.sub("", texts).strip()
    if tagged and not re.search(r"\s", tagged):
        # a single word: only a capitalised word, a number with a unit, or punctuation-ended
        w = tagged
        if not (re.fullmatch(r"[A-Z][a-z’'à-ÿ]+[.!?…:]*|[A-Z]{2,}[!?.]*|[A-Z][a-z]+[A-Z]?[a-z]*[.!?…:]", w)):
            return None
        if re.fullmatch(r"[A-Z][a-z]+[A-Z]\w*", w):
            return None  # camelCase / PascalCase identifiers
    else:
        ws = re.findall(r"\S+", tagged)
        codey = [w for w in ws if re.search(r"[-_]{1,2}[a-z]|^[a-z]+[A-Z]|^[.#]|[=;{}]|^\d+%?$", w)]
        if ws and len(codey) * 2 >= len(ws) and not re.search(r"[.!?…]$", tagged):
            return None
    if re.fullmatch(r"(\s*<[^>]*>\s*)+", texts):
        return None
    nwords = len(vis.split())
    if "“" in texts or "”" in texts or DIALOG_WHERE.search(where or "") or "(comms)" in texts:
        return "Dialog"
    if nwords <= 3 and not re.search(r"[.!?…]\s*$", vis.strip()):
        return "Label"
    return "Text"


# --------------------------------------------------------------------------
# HTML page: text between tags and a few attributes
# --------------------------------------------------------------------------
class _Html(HTMLParser):
    ATTRS = ("title", "aria-label", "placeholder", "alt")

    def __init__(self, src):
        super().__init__(convert_charrefs=True)
        self.src = src
        self.offs = [0]
        for ln in src.splitlines(keepends=True):
            self.offs.append(self.offs[-1] + len(ln))
        self.out = []  # (start, end, text, where)
        self.skip = 0
        self.script = []  # (start, end) of inline scripts
        self.stack = []

    def pos(self):
        ln, col = self.getpos()
        return self.offs[ln - 1] + col

    def handle_starttag(self, tag, attrs):
        p = self.pos()
        raw = self.get_starttag_text()
        if tag in ("script", "style", "svg"):
            self.skip += 1
        if tag == "script" and not dict(attrs).get("src"):
            self.script.append(p + len(raw))
        ident = dict(attrs).get("id") or dict(attrs).get("class", "").split(" ")[0]
        self.stack.append((tag, ident))
        if self.skip:
            return
        for name, val in attrs:
            if name in self.ATTRS and val and WORD.search(val):
                m = re.search(r'\s%s\s*=\s*(["\'])(.*?)\1' % re.escape(name), raw, re.S)
                if m:
                    self.out.append((p + m.start(2), p + m.end(2), val, "%s[%s]" % (self._where(), name), m.group(1)))
        if tag in ("meta", "link", "br", "img", "input", "use", "path", "hr", "source"):
            self.stack.pop()

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if self.stack and self.stack[-1][0] == tag:
            self.stack.pop()

    def handle_endtag(self, tag):
        if tag == "script" and self.script and isinstance(self.script[-1], int):
            self.script[-1] = (self.script[-1], self.pos())
        if tag in ("script", "style", "svg"):
            self.skip = max(0, self.skip - 1)
        while self.stack:
            t, _ = self.stack.pop()
            if t == tag:
                break

    def _where(self):
        named = [i for t, i in self.stack if i]
        return "#" + named[-1] if named else (self.stack[-1][0] if self.stack else "page")

    def handle_data(self, data):
        if self.skip or not data.strip() or not WORD.search(data):
            return
        p = self.pos()
        end = self.src.find("<", p)
        raw = self.src[p:len(self.src) if end < 0 else end]
        lead = len(raw) - len(raw.lstrip())
        raw = raw.strip()
        txt = html.unescape(raw)
        if "{" in txt or "}" in txt:
            return
        self.out.append((p + lead, p + lead + len(raw), txt, self._where(), None))


# --------------------------------------------------------------------------
# Collect every line from every file
# --------------------------------------------------------------------------
class Row:
    __slots__ = ("id", "tab", "file", "line", "kind", "where", "text", "items", "scan", "span", "html", "legend", "rng",
                 "tpl")

    def __init__(self, **k):
        for a in self.__slots__:
            setattr(self, a, k.get(a))


def line_no(src, off):
    return src.count("\n", 0, off) + 1


def js_rows(path, src, base=0, whole=None, everything=False, scan=None, names=frozenset()):
    whole = whole if whole is not None else src
    scan = scan or Scan(src)
    rows = []
    for ln in scan.lines():
        items = scan.flatten(ln)
        where = scan.where.get(ln["k"], "")
        kind = "Text" if everything else classify(scan, ln, items, where, src, names)
        if not kind:
            continue
        if any(path.endswith(f) and re.match(w, where) for f, w in SKIP):
            continue
        text = show(items)
        if len(text) > 120 and len(" ".join(visible(re.sub(r"\{[^{}]*\}", " ", text)).split())) < len(text) * 0.35:
            kind = "Markup"
        rows.append(Row(file=path, line=line_no(whole, base + scan.toks[ln["s"]].s), kind=kind, where=where,
                        text=text, items=items, scan=scan, span=base, legend=legend(items), rng=scan.extent(ln),
                        tpl=len(ln["segs"]) == 1 and ln["segs"][0][0] == "T"))
    return rows


def html_rows(path, src, everything=False, names=frozenset()):
    p = _Html(src)
    p.feed(src)
    p.close()
    rows = []
    for a, b, txt, where, q in p.out:
        kind = "Label" if len(txt.split()) <= 3 and not re.search(r"[.!?…]$", txt) else "Text"
        rows.append(Row(file=path, line=line_no(src, a), kind=kind, where=where, text=txt, html=(a, b, q)))
    for sc in p.script:
        if isinstance(sc, tuple):
            rows += js_rows(path, src[sc[0]:sc[1]], base=sc[0], whole=src, everything=everything, names=names)
    return rows


def collect(root=ROOT, files=None, everything=False):
    """All rows, in tab order, with ids. files: {path: source} to scan instead of reading disk.
    everything: every literal, not just the ones that look like text (to check an edit read back)."""
    rows, srcs, scans, names = [], {}, {}, set()
    for tab, paths in SOURCES:
        for path in paths:
            full = os.path.join(root, path)
            if files is not None and path in files:
                srcs[path] = files[path]
            elif os.path.exists(full):
                with open(full, encoding="utf-8") as f:
                    srcs[path] = f.read()
            if path in srcs and path.endswith(".js"):
                scans[path] = Scan(srcs[path])
                names |= code_words(scans[path])
    for tab, paths in SOURCES:
        for path in paths:
            if path not in srcs:
                continue
            src = srcs[path]
            got = (html_rows(path, src, everything, names) if path.endswith(".html") else
                   js_rows(path, src, everything=everything, scan=scans[path], names=names))
            got.sort(key=lambda r: r.line)
            seen = {}
            stem = os.path.splitext(os.path.basename(path))[0]
            for r in got:
                h = hashlib.sha1(r.text.encode("utf-8")).hexdigest()[:8]
                seen[h] = seen.get(h, 0) + 1
                r.id = "%s:%s%s" % (stem, h, "" if seen[h] == 1 else "~%d" % seen[h])
                r.tab = tab
                rows.append(r)
    return rows


# --------------------------------------------------------------------------
# Writing an edited line back into the source
# --------------------------------------------------------------------------
class EditError(Exception):
    pass


def split_placeholders(text):
    """'{a} is {b(c)}!' -> ['', 'a', ' is ', 'b(c)', '!'] (text, code, text, ...)."""
    out, buf, i, n = [], [], 0, len(text)
    while i < n:
        c = text[i]
        if c == "{":
            depth, j = 1, i + 1
            while j < n and depth:
                if text[j] == "{":
                    depth += 1
                elif text[j] == "}":
                    depth -= 1
                j += 1
            if depth:
                raise EditError("a { has no matching }")
            out.append("".join(buf))
            out.append(re.sub(r"\s+", " ", text[i + 1:j - 1]).strip())
            buf = []
            i = j
        elif c == "}":
            raise EditError("a } has no matching {")
        else:
            buf.append(c)
            i += 1
    out.append("".join(buf))
    return out


def js_quote(s, q):
    out = []
    for c in s:
        if c == "\\":
            out.append("\\\\")
        elif c == q:
            out.append("\\" + q)
        elif c == "\n":
            out.append("\\n")
        elif c == "\r":
            out.append("\\r")
        elif c == "\t":
            out.append("\\t")
        elif c in "  ":
            out.append("\\u%04x" % ord(c))
        else:
            out.append(c)
    return "".join(out)


def tpl_quote(s):
    s = js_quote(s, "`")
    return s.replace("${", "\\${")


def as_written(old, new):
    """The spaces at either end of a line are hard to see in a spreadsheet: they stay as they were."""
    lead = old[:len(old) - len(old.lstrip())]
    trail = old[len(old.rstrip()):]
    return lead + new.strip() + trail


def edits_for(row, new):
    """Source replacements (start, end, text) that turn row.text into new."""
    if row.html:
        a, b, q = row.html
        new = new.strip()
        if "{" in new or "}" in new:
            raise EditError("page text cannot contain { or }")
        enc = new.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if q:
            enc = enc.replace(q, "&quot;" if q == '"' else "&#39;")
        return [(a, b, enc)]
    parts = split_placeholders(as_written(row.text, new))
    want = marks(row.items)
    new_codes = parts[1::2]
    if new_codes != [m[0] for m in want]:
        return rebuild(row, parts)
    # the text between code placeholders; a slot like {target} is part of the text around it
    texts, cur = [], parts[0]
    for i, (code, kind) in enumerate(want):
        if kind == "t":
            cur += "{" + code + "}" + parts[2 * i + 2]
        else:
            texts.append(cur)
            cur = parts[2 * i + 2]
    texts.append(cur)
    # gaps: the text slots (and the expression around them) between placeholders
    gaps, cur = [], []
    xs = []
    for it in row.items:
        if it[0] == "t":
            cur.append(it)
        else:
            gaps.append(cur)
            xs.append(it)
            cur = []
    gaps.append(cur)
    scan, T, base = row.scan, row.scan.toks, row.span or 0
    out = []
    for gi, (gap, txt) in enumerate(zip(gaps, texts)):
        old = "".join(it[1] for it in gap)
        if txt == old:
            continue
        if gap:
            for n_, it in enumerate(gap):
                val = txt if n_ == 0 else ""
                kind, k = it[2]
                t = T[k]
                if kind == "s":
                    q = scan.src[t.s]
                    out.append((base + t.s, base + t.e, q + js_quote(val, q) + q))
                else:
                    out.append((base + t.rs, base + t.re_, tpl_quote(val)))
            continue
        # no literal here yet: add one next to the placeholder (only + chains have such gaps)
        lit = "'" + js_quote(txt, "'") + "'"
        if gi < len(xs) and xs[gi][2] and xs[gi][2][0] == "e":
            a = T[xs[gi][2][1]].s
            if xs[gi][2][3]:
                # values added up before the text: keep them added up, e.g. 'Day '+(a+b)
                b = T[xs[gi][2][2] - 1].e
                out.append((base + a, base + a, lit + "+("))
                out.append((base + b, base + b, ")"))
            else:
                out.append((base + a, base + a, lit + "+"))
        elif gi > 0 and xs[gi - 1][2] and xs[gi - 1][2][0] == "e":
            b = T[xs[gi - 1][2][2] - 1].e
            out.append((base + b, base + b, "+" + lit))
        else:
            raise EditError("cannot add words at that spot; put them inside a neighbouring piece of text")
    return out


def expected(row, new):
    """What resolved() should give for row once new is written."""
    parts = split_placeholders(as_written(row.text, new))
    names = tokens(row.items)
    code = {}
    for nm, it in zip(names, [x for x in row.items if x[0] == "x"]):
        code.setdefault(nm, it[1])
    return "".join(p if i % 2 == 0 else "{" + unparen(code[p]) + "}" if p in code else "{" + p + "}"
                   for i, p in enumerate(parts))


def unparen(code):
    """a?b:c and (a?b:c) are the same placeholder."""
    code = re.sub(r"\s+", " ", code).strip()
    while code.startswith("(") and code.endswith(")"):
        depth = 0
        for i, c in enumerate(code):
            depth += c == "("
            depth -= c == ")"
            if depth == 0 and i < len(code) - 1:
                return code
        code = code[1:-1].strip()
    return code


def resolved(items):
    """A line with each placeholder spelled out as its code: what the game will actually build."""
    return "".join(x[1] if x[0] == "t" else "{" + unparen(x[1]) + "}" for x in items)


def rebuild(row, parts):
    """The placeholders were dropped, repeated or moved: write the whole line again from the edited text."""
    names = tokens(row.items)
    code = {}
    for nm, it in zip(names, [x for x in row.items if x[0] == "x"]):
        code.setdefault(nm, it)
    slots = {m[0] for m in marks(row.items) if m[1] == "t"}
    seq = []
    for i, p in enumerate(parts):
        if i % 2 == 0:
            seq.append(("t", p))
        elif p in code:
            seq.append(("x", code[p]))
        elif p in slots:
            seq.append(("t", "{" + p + "}"))
        else:
            ok = " ".join("{%s}" % n for n in dict.fromkeys(names + sorted(slots)))
            raise EditError("{%s} is not one of this line's placeholders (%s)" % (p, ok or "it has none"))
    merged = []
    for kind, v in seq:
        if kind == "t" and merged and merged[-1][0] == "t":
            merged[-1] = ("t", merged[-1][1] + v)
        elif not (kind == "t" and v == ""):
            merged.append((kind, v))
    src = row.scan.src
    if row.tpl:
        body = "".join(tpl_quote(v) if k == "t" else "${" + v[3] + "}" for k, v in merged)
        out = "`" + body + "`"
    else:
        strs = [x for x in row.items if x[0] == "t" and x[2][0] == "s"]
        q = src[row.scan.toks[strs[0][2][1]].s] if strs else "'"
        pieces = []
        for k, v in merged:
            if k == "t":
                pieces.append(q + js_quote(v, q) + q)
            elif v[4] == "op" or SIMPLE.fullmatch(v[3].strip()):
                pieces.append(v[3])
            else:
                pieces.append("(" + v[3] + ")")
        # two values side by side at the start would be added up as numbers: start from a string
        if not merged or merged[0][0] == "x" and (len(merged) == 1 or merged[1][0] == "x"):
            pieces.insert(0, q + q)
        out = "+".join(pieces)
    a, b = row.rng
    base = row.span or 0
    return [(base + a, base + b, out)]


def node_check(path, src):
    node = shutil.which("node")
    if not node or not path.endswith(".js"):
        return None
    with tempfile.NamedTemporaryFile("w", suffix=".js", delete=False, encoding="utf-8") as f:
        f.write(src)
        tmp = f.name
    try:
        r = subprocess.run([node, "--check", tmp], capture_output=True, text=True)
        return None if r.returncode == 0 else r.stderr.strip().splitlines()[-1] if r.stderr.strip() else "syntax error"
    finally:
        os.unlink(tmp)


def db_names(root=ROOT):
    """Every piece of text in the game database, with the table and column it sits in."""
    import json
    out = {}
    try:
        with open(os.path.join(root, "game/data/db.json"), encoding="utf-8") as f:
            db = json.load(f)
    except (OSError, ValueError):
        return out
    for table, rws in db.items():
        for r in rws if isinstance(rws, list) else []:
            for col, v in (r.items() if isinstance(r, dict) else []):
                if isinstance(v, str) and v.strip():
                    out.setdefault(v.strip(), "%s.%s" % (table, col))
    return out


def twins(changes, root=ROOT):
    """Warnings: an edited name that is written somewhere else too, unedited. The game may match the two (a building's
    name, a ship's name in the database), so changing one alone can break that."""
    rows = collect(root)
    by_id = {r.id: r for r in rows}
    every = collect(root, everything=True)
    db = db_names(root)
    out = []
    for rid in changes:
        r = by_id.get(rid)
        if r is None or "{" in r.text or len(r.text) > 60:
            continue
        key = r.text.strip()
        also = ["%s:%d" % (x.file, x.line) for x in every
                if x.text.strip().lower() == key.lower() and (x.file, x.line) != (r.file, r.line) and x.id not in changes
                and not (x.file == r.file and x.items and r.items and x.items[0][2] == r.items[0][2])]
        if len(also) > 3:
            also = []   # a word used all over (Continue, Cancel) is a common label, not a name the code matches
        if key in db:
            also.append("the game database (%s)" % db[key])
        if also:
            out.append("%s (%s line %d): \u201c%s\u201d is also written at %s. If the game matches these up, change "
                       "them together." % (rid, r.file, r.line, key, ", ".join(also[:6]) + (" and more" if len(also) > 6 else "")))
    return out


def apply_edits(changes, root=ROOT, write=True):
    """changes: {id: new text}. Returns (applied [(row, new)], problems [str], new sources {path: src})."""
    rows = collect(root)
    by_id = {r.id: r for r in rows}
    per_file, applied, problems = {}, [], []
    for rid, new in changes.items():
        r = by_id.get(rid)
        if r is None:
            problems.append("%s: this line is no longer in the code (it was changed since the export). Export again and redo this edit." % rid)
            continue
        try:
            per_file.setdefault(r.file, []).extend(edits_for(r, new))
            applied.append((r, new))
        except EditError as ex:
            problems.append("%s (%s line %d, %s): %s" % (rid, r.file, r.line, r.where or r.kind, ex))
    if problems:
        return applied, problems, {}
    out = {}
    for path, eds in per_file.items():
        full = os.path.join(root, path)
        with open(full, encoding="utf-8") as f:
            src = f.read()
        eds.sort(key=lambda e: (e[0], e[1]))
        for a, b in zip(eds, eds[1:]):
            if b[0] < a[1]:
                problems.append("%s line %d: two edited lines sit inside each other (one is a placeholder of the "
                                "other). Import one now and the other in a second pass." % (path, line_no(src, b[0])))
        for a, b, txt in reversed(eds):
            src = src[:a] + txt + src[b:]
        try:
            if path.endswith(".js"):
                lex(src)
        except LexError as ex:
            problems.append("%s: the edit broke the file (%s)" % (path, ex))
        err = node_check(path, src)
        if err:
            problems.append("%s: the edit broke the file: %s" % (path, err))
        out[path] = src
    if problems:
        return applied, problems, {}
    # every edited line must read back as written
    every = collect(root, files=out, everything=True)
    after = {r.text for r in every} | {resolved(r.items) for r in every if r.items}
    for r, new in applied:
        # as shown (a placeholder's own text may have been edited too), or as built (placeholders were moved)
        want = new.strip() if r.html else as_written(r.text, new)
        if want not in after and (r.html or expected(r, new) not in after):
            if want.strip():
                problems.append("%s: after the edit the line does not read back as written; nothing was saved" % r.id)
    if problems:
        return applied, problems, {}
    if write:
        for path, src in out.items():
            with open(os.path.join(root, path), "w", encoding="utf-8", newline="") as f:
                f.write(src)
    return applied, problems, out


# --------------------------------------------------------------------------
# Workbook
# --------------------------------------------------------------------------
HEAD = ["Kind", "Where", "Text", "Original", "Placeholders", "File", "Line", "ID"]
WIDTH = {"Kind": 9, "Where": 34, "Text": 80, "Original": 50, "Placeholders": 40, "File": 18, "Line": 7, "ID": 18}
GUIDE = [
    ("STAR REBELLION: ALL THE GAME'S TEXT", "title"),
    ("Every line of dialog, briefing, tooltip and label the game shows, pulled out of the code. Edit the Text column, "
     "download the workbook as .xlsx and send it back (or run `python3 tools/text/text.py import` on it). The tool writes "
     "your words back into the code.", "text"),
    ("", "text"),
    ("HOW TO EDIT", "h"),
    ("1. Change only the Text column (white). Changed rows turn yellow so you can see what you touched.", "text"),
    ("2. Original (grey) is what the game says now. File, Line and ID say where it lives. They are ignored on import, "
     "but do not delete the ID column or reorder rows across tabs: the ID is how the tool finds the line.", "text"),
    ("3. {Curly braces} are filled in by the game: {t.first} is a rebel's first name, {target} the mission's target, "
     "{n} a number, and so on. A long piece of code shows as a number, {1}, and the Placeholders column says what it "
     "is. Keep every placeholder, exactly as written and in the same order. Move words around them freely.", "text"),
    ("4. Some text carries HTML: <b>bold</b>, <br> for a line break, <span class=...> for colour. Keep the tags around "
     "the words they wrap.", "text"),
    ("5. Kind: Dialog is someone talking (barks, comms, briefings, the merchant). Text is narration and descriptions. "
     "Label is a short button, title or tag; check it still fits before making it longer.", "text"),
    ("6. Use the filter on row 1 to show one Kind, or search (Ctrl+F) for a line you saw in the game.", "text"),
    ("7. To ask for a change instead of making it (a value that should not be fixed, a line to delete), write "
     "[Note for Claude: ...] or [Remove] in the Text cell. Those rows are never written into the game; the import "
     "lists them for Claude to act on.", "text"),
    ("8. Sheets reads a cell that starts with + or = as a formula. To start a line with +, type an apostrophe first: "
     "'+5% mission success.", "text"),
    ("", "text"),
    ("WHAT IS NOT HERE", "h"),
    ("Item, ship, weapon and enemy names and descriptions live in the game database (python3 tools/db/build.py "
     "export-xlsx). Rebel names come from lists in game/js/rebel.js and are left out on purpose. A line the game builds "
     "from many small pieces of code may be missing too; ask for it by quoting what you see on screen.", "text"),
    ("", "text"),
    ("IF THE IMPORT REFUSES", "h"),
    ("It names the row and the reason (a placeholder went missing, or the code changed that line since this export) and "
     "saves nothing. Fix the row, or export a fresh workbook and copy your edits across.", "text"),
]


def _need_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        sys.exit("openpyxl is required: pip install openpyxl")


def export_xlsx(rows, out):
    _need_openpyxl()
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as L

    def F(**k):
        k.setdefault("size", 10)
        return Font(name="Arial", **k)

    head_fill = PatternFill("solid", fgColor="1F2A44")
    grey = PatternFill("solid", fgColor="EFEFEF")
    yellow = PatternFill("solid", fgColor="FFF2B3")
    wb = Workbook()
    g = wb.active
    g.title = "Guide"
    g.column_dimensions["A"].width = 120
    for i, (txt, style) in enumerate(GUIDE, start=1):
        c = g.cell(row=i, column=1, value=txt)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = F(bold=True, size=14) if style == "title" else F(bold=True) if style == "h" else F()
    g.sheet_view.showGridLines = False
    counts = {}
    for tab, _ in SOURCES:
        rs = [r for r in rows if r.tab == tab]
        if not rs:
            continue
        ws = wb.create_sheet(tab)
        for j, h in enumerate(HEAD, start=1):
            c = ws.cell(row=1, column=j, value=h)
            c.font = F(bold=True, color="FFFFFF")
            c.fill = head_fill
            ws.column_dimensions[L(j)].width = WIDTH[h]
        for i, r in enumerate(rs, start=2):
            vals = {"Kind": r.kind, "Where": r.where, "Text": r.text, "Original": r.text, "Placeholders": r.legend or None,
                    "File": r.file, "Line": r.line, "ID": r.id}
            for j, h in enumerate(HEAD, start=1):
                c = ws.cell(row=i, column=j, value=vals[h])
                c.alignment = Alignment(wrap_text=h in ("Text", "Original", "Where", "Placeholders"), vertical="top")
                if h == "Text":
                    c.font = F(size=11)
                else:
                    c.font = F(color="666666")
                    c.fill = grey
        last = len(rs) + 1
        tcol, ocol = L(HEAD.index("Text") + 1), L(HEAD.index("Original") + 1)
        ws.conditional_formatting.add("%s2:%s%d" % (tcol, tcol, last),
                                      FormulaRule(formula=["EXACT(%s2,%s2)=FALSE" % (tcol, ocol)], fill=yellow))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = "A1:%s%d" % (L(len(HEAD)), last)
        counts[tab] = len(rs)
    wb.save(out)
    return counts


def read_xlsx(path):
    """{id: text} for every row, plus {id: original}."""
    _need_openpyxl()
    from openpyxl import load_workbook
    # formulas as written, not their results: Sheets turns a line typed as "+5% ..." into the formula "=5% ..."
    wb = load_workbook(path, data_only=False)
    texts, origs = {}, {}
    for ws in wb.worksheets:
        head = [c.value for c in ws[1]]
        if "ID" not in head or "Text" not in head:
            continue
        ci, ct = head.index("ID"), head.index("Text")
        co = head.index("Original") if "Original" in head else None
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or ci >= len(row) or not row[ci]:
                continue
            rid = str(row[ci]).strip()
            v = row[ct]
            o = row[co] if co is not None else None
            if isinstance(v, str) and v.startswith("=") and not (isinstance(o, str) and o.startswith("=")):
                v = (o[0] if isinstance(o, str) and o[:1] in "+-" else "") + v[1:]
            texts[rid] = "" if v is None else str(v).replace("\r\n", "\n").replace("\r", "\n")
            if co is not None:
                o = row[co]
                origs[rid] = "" if o is None else str(o).replace("\r\n", "\n").replace("\r", "\n")
    return texts, origs


NOTE = re.compile(r"\[\s*(note for claude|remove|delete|deprecate)\b", re.I)


def notes(changes):
    """Rows that ask Claude for something ([Note for Claude: ...], [Remove]) rather than give the new words.
    They are never written into the game: they are listed for Claude to act on by hand."""
    return {rid: t for rid, t in changes.items() if NOTE.search(t)}


def changed(texts, origs, rows):
    """The rows the writer actually changed: text differs from what the code says now (or from the exported original)."""
    by_id = {r.id: r for r in rows}
    out = {}
    for rid, txt in texts.items():
        r = by_id.get(rid)
        base = r.text if r is not None else origs.get(rid)
        if base is None:
            continue
        if txt.strip() == base.strip():
            continue
        if r is None and txt.strip() == (origs.get(rid) or "").strip():
            continue
        out[rid] = txt
    return out


# --------------------------------------------------------------------------
KINDS = ("Dialog", "Text", "Label", "Markup")


def selftest():
    """Edit every line at once in a copy of the game, write it back, and check every line reads back as written,
    every file still parses, and an export/import through a workbook changes exactly the edited line."""
    tmp = tempfile.mkdtemp(prefix="sr-text-")
    fails = []
    try:
        for _, paths in SOURCES:
            for path in paths:
                os.makedirs(os.path.dirname(os.path.join(tmp, path)), exist_ok=True)
                shutil.copy(os.path.join(ROOT, path), os.path.join(tmp, path))
        rows = collect(tmp)
        before = collect(tmp, everything=True)
        ids = [r.id for r in rows]
        if len(set(ids)) != len(ids):
            fails.append("duplicate ids")
        # every line: words added at the start, in the middle and at the end, and a quote, backslash and newline
        def mangle(r):
            parts = re.split(r"(\{[^{}]*\})", r.text)
            parts[0] = "Wq " + parts[0]
            parts[-1] = parts[-1] + " zap"
            if len(parts) > 2:
                parts[2] = " mid’s 'x' \"y\" \\ $z `t`\n" + parts[2]
            return "".join(parts)
        js = {r.id: mangle(r) for r in rows if not r.html}
        page = {r.id: "Wq " + r.text + " & <zap> \"" for r in rows if r.html}
        applied, problems, out = apply_edits(dict(js, **page), root=tmp)
        fails += problems[:20]
        if not problems:
            after = collect(tmp, everything=True)
            got = {r.text for r in after}
            by_id = {r.id: r for r in rows}
            miss = [rid for rid, t in js.items() if as_written(by_id[rid].text, t) not in got]
            miss += [rid for rid, t in page.items() if t.strip() not in got]
            fails += ["%s does not read back after an edit" % m for m in miss[:20]]
            if len(after) != len(before):
                fails.append("%d literals before the edit, %d after" % (len(before), len(after)))
        # placeholders moved, repeated and dropped: the line is rebuilt (on a fresh copy)
        for _, paths in SOURCES:
            for path in paths:
                shutil.copy(os.path.join(ROOT, path), os.path.join(tmp, path))
        rows = collect(tmp)
        moved, taken = {}, {}
        for r in rows:
            names = [m[0] for m in marks(r.items) if m[1] == "x"] if r.items else []
            if not names or any(a < r.rng[1] and r.rng[0] < b for a, b in taken.get(r.file, [])):
                continue
            taken.setdefault(r.file, []).append(r.rng)
            back = list(reversed(names))
            moved[r.id] = "Wq " + " then ".join("{%s}" % n for n in back + back[:1]) + " zap"
        _, problems, _ = apply_edits(moved, root=tmp)
        fails += problems[:20]
        # a placeholder that is not the line's own is refused, and nothing is written
        r = next(r for r in rows if "{" in r.text and not r.html)
        _, problems, _ = apply_edits({r.id: r.text + " {nope}"}, root=tmp, write=False)
        if not problems:
            fails.append("an unknown placeholder was accepted")
        # workbook round trip on a fresh copy: one edit in, one line changed
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            print("openpyxl missing: workbook round trip skipped")
        else:
            for _, paths in SOURCES:
                for path in paths:
                    shutil.copy(os.path.join(ROOT, path), os.path.join(tmp, path))
            rows = collect(tmp)
            xl = os.path.join(tmp, "t.xlsx")
            export_xlsx(rows, xl)
            from openpyxl import load_workbook
            wb = load_workbook(xl)
            ws = wb["Ground"]
            head = [c.value for c in ws[1]]
            cell = ws.cell(row=5, column=head.index("Text") + 1)
            target = ws.cell(row=5, column=head.index("ID") + 1).value
            cell.value = "Zq " + cell.value
            wb.save(xl)
            texts, origs = read_xlsx(xl)
            ch = changed(texts, origs, rows)
            if list(ch) != [target]:
                fails.append("workbook round trip found %d edits, expected 1" % len(ch))
            else:
                _, problems, out = apply_edits(ch, root=tmp)
                fails += problems
                after = {r.id for r in collect(tmp)}
                lost = set(r.id for r in rows) - after
                if len(lost) != 1:
                    fails.append("workbook round trip changed %d lines, expected 1" % len(lost))
        print("text selftest: %d lines, %d rebuilt, %s" % (len(ids), len(moved), "ok" if not fails else "%d problem(s)" % len(fails)))
        for f in fails:
            print("  " + f)
        return 1 if fails else 0
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("export", help="write every line of game text to a workbook")
    e.add_argument("out", nargs="?", default="star-rebellion-text.xlsx")
    i = sub.add_parser("import", help="write the edited Text column back into the code")
    i.add_argument("path")
    i.add_argument("--dry-run", action="store_true")
    f = sub.add_parser("find", help="where does a line of text live")
    f.add_argument("words", nargs="+")
    sub.add_parser("report", help="how many lines per tab and kind")
    sub.add_parser("dump", help="every line, tab separated (for scripts and diffs)")
    sub.add_parser("spans", help=argparse.SUPPRESS)      # for tools/text-smoke.js
    sub.add_parser("selftest", help=argparse.SUPPRESS)   # for tools/text-smoke.js
    a = ap.parse_args()

    if a.cmd == "export":
        counts = export_xlsx(collect(), a.out)
        print("Wrote %s: %d lines (%s)" % (a.out, sum(counts.values()), ", ".join("%s %d" % kv for kv in counts.items())))
    elif a.cmd == "import":
        texts, origs = read_xlsx(a.path)
        rows = collect()
        ch = changed(texts, origs, rows)
        held = notes(ch)
        ch = {k: v for k, v in ch.items() if k not in held}
        by_id = {r.id: r for r in rows}
        if held:
            print("Notes for Claude (not imported; these need a code change, so ask Claude to act on them):")
            for rid, t in held.items():
                r = by_id.get(rid)
                where = "%s line %d  [%s]" % (r.file, r.line, r.where or r.kind) if r else rid
                print("  %s\n    now: %s\n    note: %s" % (where, (r.text if r else "?").replace("\n", "\\n"), t.replace("\n", "\\n")))
            print()
        if not ch:
            print("No other edits found in %s." % a.path)
            return
        applied, problems, _ = apply_edits(ch, write=not a.dry_run)
        by_id = {r.id: r for r in rows}
        for rid, new in ch.items():
            r = by_id.get(rid)
            if r is None:
                continue
            print("%s line %d  [%s]\n  - %s\n  + %s" % (r.file, r.line, r.where or r.kind, r.text.replace("\n", "\\n"),
                                                        new.strip().replace("\n", "\\n")))
        if problems:
            print("\nNothing was saved. Fix these rows and import again:", file=sys.stderr)
            for p in problems:
                print("  " + p, file=sys.stderr)
            sys.exit(1)
        warn = twins(ch)
        if warn:
            print("\nCheck these (saved or not, they may need a matching change elsewhere):")
            for w in warn:
                print("  " + w)
        if a.dry_run:
            print("\n%d edit(s) would be written. Run again without --dry-run to save them." % len(applied))
        else:
            files = sorted({r.file for r, _ in applied})
            print("\nSaved %d edit(s) to %s." % (len(applied), ", ".join(files)))
            if any(f.startswith("game/data/") for f in files):
                print("Run python3 tools/db/build.py export-js to refresh db.js.")
    elif a.cmd == "find":
        q = " ".join(a.words).lower()
        hits = [r for r in collect() if q in r.text.lower() or q in (r.where or "").lower()]
        for r in hits:
            print("%s:%d  [%s %s]  %s" % (r.file, r.line, r.kind, r.where, r.text.replace("\n", "\\n")))
        print("%d match(es)" % len(hits))
    elif a.cmd == "report":
        rows = collect()
        print("%-8s %6s %6s %6s %6s %6s" % ("Tab", "Dialog", "Text", "Label", "Markup", "All"))
        for tab, _ in SOURCES:
            rs = [r for r in rows if r.tab == tab]
            k = {x: sum(1 for r in rs if r.kind == x) for x in KINDS}
            print("%-8s %6d %6d %6d %6d %6d" % (tab, k["Dialog"], k["Text"], k["Label"], k["Markup"], len(rs)))
        print("%-8s %34d" % ("All", len(rows)))
    elif a.cmd == "spans":
        import json
        out = {}
        for _, paths in SOURCES:
            for path in paths:
                if not path.endswith(".js"):
                    continue
                with open(os.path.join(ROOT, path), encoding="utf-8") as fh:
                    src = fh.read()
                u16, n = [], 0
                for ch in src:
                    u16.append(n)
                    n += 2 if ord(ch) > 0xFFFF else 1
                u16.append(n)
                out[path] = [[u16[t.s], u16[t.e], t.v] for t in lex(src) if t.t in ("str", "tpl", "tpl_head", "tpl_mid", "tpl_tail")]
        print(json.dumps(out))
    elif a.cmd == "selftest":
        sys.exit(selftest())
    elif a.cmd == "dump":
        for r in collect():
            print("\t".join([r.id, r.kind, r.file, str(r.line), r.where or "", r.text.replace("\n", "\\n")]))


if __name__ == "__main__":
    main()
